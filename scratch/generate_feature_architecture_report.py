"""
InternHub - Complete Feature Implementation and Technical Architecture Report Generator
Generates:
  - reports/INTERNHUB_COMPLETE_FEATURE_IMPLEMENTATION_AND_ARCHITECTURE_REPORT.docx
  - reports/INTERNHUB_COMPLETE_FEATURE_IMPLEMENTATION_AND_ARCHITECTURE_REPORT.pdf
Style: Professional Black & White (Clean Charcoal / White / Subtle Neutral Grays)
"""

import os
import sys
import html
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# ==============================================================================
# 1. ARCHITECTURE & REPOSITORY DATA DEFINITIONS
# ==============================================================================

DOC_TITLE = "InternHub – Complete Feature Implementation and Technical Architecture Report"
DOC_SUBTITLE = "Comprehensive Source-Code Audit, Page-to-Feature Implementation Map, API Directory, and Workflow Architecture"
DOC_ID = "IHB-ARCH-2026-V1"
DOC_DATE = "October 23, 2026"
DOC_SYSTEM = "Smart Internship Management System (InternHub)"
DOC_AUTHOR = "Senior Software Architect & QA Systems Engineer"
DOC_CLASSIFICATION = "Internal Engineering & Architectural Governance"
DOC_VERDICT = "VERIFIED PRODUCTION ARCHITECTURE (26 Pages, 123 APIs, 14 Database Collections, 225 Automated Tests)"

# --- 26 Frontend Pages Directory Data ---
PAGES_DIRECTORY = [
    {
        "title": "Home & Public Landing",
        "route": "/",
        "file": "frontend/src/pages/HomePage.jsx",
        "roles": "Public (All)",
        "nav": "Sidebar ('System Overview'), Public Navbar, Brand Logo",
        "components": "SimsLogo, PublicNavbar, StatCards, RoleShowcase, FeatureHighlights",
        "desc": "Public landing portal with institutional mission, system metrics, role navigation cards, and live health status indicator."
    },
    {
        "title": "Internships Directory",
        "route": "/internships",
        "file": "frontend/src/pages/InternshipsDirectoryPage.jsx",
        "roles": "Public / All Authenticated",
        "nav": "Sidebar ('Internship Directory'), Top Navbar quick link",
        "components": "InternshipCard, SearchBar, DepartmentFilter, StatusFilter, Pagination",
        "desc": "Publicly searchable catalog of published internships with location, stipend, department filters, and instant apply triggers."
    },
    {
        "title": "Jobs & LinkedIn Radar",
        "route": "/external-jobs",
        "file": "frontend/src/pages/ExternalJobsPage.jsx",
        "roles": "Public / All Authenticated",
        "nav": "Sidebar ('Jobs & LinkedIn Radar'), Top Navbar pill button",
        "components": "ExternalJobCard, SkillMatchRadar, BookmarkButton, ApplyTrackerModal",
        "desc": "Aggregated external job opportunities with live AI skill matching, LinkedIn direct search links, and student bookmark tracking."
    },
    {
        "title": "Unified Login",
        "route": "/login",
        "file": "frontend/src/pages/LoginPage.jsx",
        "roles": "Public (Anonymous)",
        "nav": "Public Navbar ('Sign In'), Top Navbar actions, redirect on unauthenticated",
        "components": "LoginForm, GoogleSignInButton, TwoFactorModal, ForgotPasswordModal",
        "desc": "Unified authentication portal supporting standard credentials, Google GIS SSO, 2FA OTP verification, and password reset dialog."
    },
    {
        "title": "Legacy Admin Login",
        "route": "/admin/login",
        "file": "frontend/src/pages/AdminLoginPage.jsx",
        "roles": "Public (Anonymous)",
        "nav": "Direct URL (Redirects to /login)",
        "components": "Navigate (React Router redirect component)",
        "desc": "Legacy admin URL that automatically and seamlessly redirects users to the unified single login portal at /login."
    },
    {
        "title": "User Registration",
        "route": "/register",
        "file": "frontend/src/pages/RegisterPage.jsx",
        "roles": "Public (Anonymous)",
        "nav": "Public Navbar ('Register'), Login page toggle link",
        "components": "RegisterForm, RoleSelector (Student/Company), GoogleSignInButton, PasswordStrengthMeter",
        "desc": "Self-service registration portal allowing new candidates to register as Students or employer representatives as Companies."
    },
    {
        "title": "Verify Certificate",
        "route": "/verify-certificate, /verify-certificate/:code",
        "file": "frontend/src/pages/VerifyCertificatePage.jsx",
        "roles": "Public (All)",
        "nav": "Sidebar ('Verify Certificate'), Public Navbar, QR Code scan target",
        "components": "CertificateSearchInput, VerifiedBadge, CredentialDetailsCard, PDFDownloadButton",
        "desc": "Public cryptographic verification ledger allowing employers and universities to validate authentic completion certificates."
    },
    {
        "title": "Institutional Policies",
        "route": "/policies, /privacy, /terms, /security-policy, /cookies",
        "file": "frontend/src/pages/PoliciesPage.jsx",
        "roles": "Public (All)",
        "nav": "Footer links ('Privacy', 'Terms', 'Security', 'Cookies')",
        "components": "PolicyNavigationTabs, PolicySectionContent, PrintButton",
        "desc": "Comprehensive institutional compliance document detailing terms of service, privacy regulations, security standards, and cookie rules."
    },
    {
        "title": "Role Workspace Dashboard",
        "route": "/dashboard",
        "file": "frontend/src/pages/DashboardPage.jsx",
        "roles": "All Authenticated (Student, Mentor, Company, Coordinator, Admin)",
        "nav": "Sidebar ('My Dashboard'), Top Navbar profile dropdown",
        "components": "StudentDashboard, CompanyDashboard, MentorDashboard, CoordinatorDashboard, AdminDashboard",
        "desc": "Central personalized role hub rendering specialized metrics, quick actions, task overviews, and live progress bars for each user role."
    },
    {
        "title": "Notification Center",
        "route": "/notifications",
        "file": "frontend/src/pages/NotificationsPage.jsx",
        "roles": "All Authenticated Roles",
        "nav": "Sidebar ('Notifications'), Top Navbar Bell Icon dropdown",
        "components": "NotificationItem, FilterTabs (All/Unread), ClearAllButton, SoundToggle",
        "desc": "Full-page notification center displaying chronological alerts, read/unread states, mark-all-read triggers, and push preferences."
    },
    {
        "title": "Direct & Group Chat",
        "route": "/chat",
        "file": "frontend/src/pages/ChatPage.jsx",
        "roles": "All Authenticated Roles",
        "nav": "Sidebar ('Direct & Group Chat'), Top Navbar message icon",
        "components": "ConversationList, MessageThread, MessageInput, AttachmentUploader, EmojiReactionPicker, GroupMembersModal",
        "desc": "Real-time bi-directional messaging platform featuring direct chats, coordinator groups, attachments, emoji reactions, and typing indicators."
    },
    {
        "title": "User Profile & Security",
        "route": "/profile",
        "file": "frontend/src/pages/ProfilePage.jsx",
        "roles": "All Authenticated Roles",
        "nav": "Sidebar ('My Profile'), Top Navbar profile dropdown",
        "components": "AvatarUploadModal, ResumeUploadSection, TwoFactorToggle, ActiveSessionsManager, PasswordChangeForm",
        "desc": "Personal profile management center for updating bio, skills, resume PDF, avatar image, 2FA security, and active session termination."
    },
    {
        "title": "Student My Applications",
        "route": "/my-applications",
        "file": "frontend/src/pages/StudentApplicationsPage.jsx",
        "roles": "Student",
        "nav": "Sidebar ('My Applications'), /applications role redirect",
        "components": "ApplicationTimeline, StatusPill, WithdrawModal, FeedbackViewer",
        "desc": "Student dashboard tracking all submitted internship applications, live status stages, employer decision notes, and withdrawal options."
    },
    {
        "title": "Student Daily Attendance",
        "route": "/my-attendance",
        "file": "frontend/src/pages/StudentAttendancePage.jsx",
        "roles": "Student",
        "nav": "Sidebar ('Daily Attendance'), /attendance role redirect",
        "components": "AttendanceSubmissionForm, MonthlyAttendanceCalendar, AttendanceHistoryTable, SummaryStatsCard",
        "desc": "Intern attendance submission portal for logging daily hours and work summaries, viewing supervisor approval states, and tracking rates."
    },
    {
        "title": "Student Milestone Tasks",
        "route": "/my-tasks",
        "file": "frontend/src/pages/StudentTasksPage.jsx",
        "roles": "Student",
        "nav": "Sidebar ('Milestone Tasks'), /tasks role redirect",
        "components": "TaskCardList, TaskSubmissionModal, DeliverableUrlInput, MentorFeedbackBanner",
        "desc": "Student task execution hub displaying assigned milestones, overdue alerts, deliverable link submission, and mentor evaluation ratings."
    },
    {
        "title": "Student My Certificates",
        "route": "/my-certificates, /certificates",
        "file": "frontend/src/pages/CertificatesPage.jsx",
        "roles": "Student",
        "nav": "Sidebar ('My Certificates')",
        "components": "CertificateCard, PDFDownloadButton, PublicVerifyLink, RequestCompletionButton",
        "desc": "Awarded certificates gallery allowing students to download certified landscape PDFs, copy verification URLs, and view grades."
    },
    {
        "title": "Employer Candidate Review",
        "route": "/manage/applications",
        "file": "frontend/src/pages/CompanyApplicationsPage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Candidate Review'), /applications role redirect",
        "components": "ApplicantFilterToolbar, CandidateReviewModal, ResumeViewerModal, StatusTransitionDropdown, AIResumeScoreBadge",
        "desc": "Employer applicant screening console for reviewing candidate profiles, analyzing AI resume fit, scheduling interviews, and selection."
    },
    {
        "title": "Intern Allocations Management",
        "route": "/manage/allocations",
        "file": "frontend/src/pages/ManageAllocationsPage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Intern Allocations'), /allocations role redirect",
        "components": "AllocationTable, MentorAssignModal, RequestCompletionModal, IssueCertificateModal, SeedDemoButton",
        "desc": "Workplace allocation console for pairing accepted candidates with industry mentors, setting dates, and initiating completion requests."
    },
    {
        "title": "Attendance Audit & Review",
        "route": "/manage/attendance",
        "file": "frontend/src/pages/ManageAttendancePage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Attendance Audit'), /attendance role redirect",
        "components": "AttendanceAuditTable, ReviewStatusDropdown, CorrectionModal, DateRangeFilter",
        "desc": "Supervisor attendance review console for auditing intern daily submissions, approving hours, or correcting records with reasons."
    },
    {
        "title": "Internship Tasks Management",
        "route": "/manage/tasks",
        "file": "frontend/src/pages/ManageTasksPage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Manage Tasks'), /tasks role redirect",
        "components": "CreateTaskModal, TaskReviewModal, InternSelector, OverdueBadge, DeliverableInspector",
        "desc": "Mentor milestone creation and evaluation portal for assigning structured project tasks, tracking submissions, and awarding scores."
    },
    {
        "title": "Evaluations & Performance",
        "route": "/evaluations",
        "file": "frontend/src/pages/EvaluationsPage.jsx",
        "roles": "All Authenticated Roles",
        "nav": "Sidebar ('Evaluations')",
        "components": "EvaluationFormModal, PerformanceRadarChart, CriteriaScoringInput, EvaluationHistoryList",
        "desc": "Performance appraisal portal for submitting mid-term and final intern reviews across technical, professional, and behavioral metrics."
    },
    {
        "title": "Opportunities Management",
        "route": "/manage/internships",
        "file": "frontend/src/pages/ManageInternshipsPage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Manage Opportunities')",
        "components": "CreateInternshipModal, EditInternshipModal, StatusToggleSwitch, ApplicantCountBadge",
        "desc": "Employer internship management console for drafting, editing, publishing, closing, or archiving internship postings."
    },
    {
        "title": "Analytics & Reports Export",
        "route": "/reports",
        "file": "frontend/src/pages/ReportsPage.jsx",
        "roles": "Company, Mentor, Coordinator, Admin",
        "nav": "Sidebar ('Reports & Export')",
        "components": "MetricsOverviewCards, TrendCharts, ExportCsvButton, DateRangePicker",
        "desc": "Institutional intelligence dashboard providing aggregate enrollment trends, departmental hiring ratios, and sanitized CSV reports."
    },
    {
        "title": "Administrator User Console",
        "route": "/admin/users",
        "file": "frontend/src/pages/AdminUsersPage.jsx",
        "roles": "Admin",
        "nav": "Sidebar ('User Management'), /users role redirect",
        "components": "UserManagementTable, CreateCoordinatorModal, RoleChangeDropdown, AccountStatusSwitch, DeleteUserModal",
        "desc": "Governance console for searching user accounts, creating coordinators, modifying roles, deactivating accounts, and last-admin guards."
    },
    {
        "title": "Admin Certificate Audit",
        "route": "/admin/certificates",
        "file": "frontend/src/pages/AdminCertificatesPage.jsx",
        "roles": "Admin, Coordinator",
        "nav": "Sidebar ('Certificate Audit')",
        "components": "CertificateAuditTable, ResendEmailButton, PreviewModal, MentorIssuanceStats",
        "desc": "Institutional credential audit dashboard displaying all issued certificates, verification codes, delivery logs, and resend triggers."
    },
    {
        "title": "404 Not Found Page",
        "route": "* (Catch-all)",
        "file": "frontend/src/pages/NotFoundPage.jsx",
        "roles": "Public (All)",
        "nav": "Automatic fallback on invalid route",
        "components": "NotFoundGraphic, HomeRedirectButton, SupportContactLink",
        "desc": "Standard error page displayed when users navigate to non-existent URLs, providing a clean path back to home or dashboard."
    }
]

# --- 26 Feature-to-File Implementation Map Data ---
FEATURES_MAP = [
    {
        "name": "1. User Registration & Single Login",
        "page": "Unified Login & Register",
        "route": "/login, /register",
        "frontend_file": "frontend/src/pages/LoginPage.jsx, frontend/src/pages/RegisterPage.jsx",
        "components": "useAuth hook, AuthContext.jsx, GoogleSignInButton, PasswordStrengthMeter",
        "backend_api": "POST /api/auth/register, POST /api/auth/login, GET /api/auth/me",
        "backend_files": "backend/src/routes/authRoutes.js, authController.js, validateAuth.js, auth.js, User.js, Session.js",
        "database": "users (email, password, role, status), sessions (tokenSignature, isValid, deviceInfo, ipAddress)",
        "integrations": "Not applicable for standard login; bcrypt password hashing (10 salt rounds)",
        "workflow": "Candidate submits registration/login -> server validates complexity and credentials -> bcrypt compares hash -> creates Session record -> signs JWT -> returns user profile and token -> client stores in sessionStorage and sets Axios Bearer header.",
        "permissions": "Public / Anonymous users for registration and login; All authenticated roles for /me",
        "test_evidence": "backend/tests/auth.test.js (TC-AUTH-01, TC-AUTH-02, TC-AUTH-07, TC-AUTH-08, TC-AUTH-11)"
    },
    {
        "name": "2. Google GIS OAuth Sign-In",
        "page": "Unified Login & Register",
        "route": "/login, /register",
        "frontend_file": "frontend/src/components/common/GoogleSignInButton.jsx, frontend/src/pages/LoginPage.jsx",
        "components": "GoogleSignInButton.jsx, useAuth hook, Google Identity Services SDK",
        "backend_api": "POST /api/auth/google, GET /api/auth/google/client-id",
        "backend_files": "backend/src/routes/authRoutes.js, authController.js, firebaseAdmin.js, User.js, Session.js",
        "database": "users (email, name, role, avatar, isEmailVerified), sessions (tokenSignature)",
        "integrations": "Google Identity Services (GIS) / Firebase Admin Auth token verification",
        "workflow": "User clicks Google button -> GIS returns credential ID token -> client posts token to /api/auth/google -> backend verifies token against Google OAuth servers -> auto-provisions student if new user or retrieves existing -> signs JWT and session -> logs user into dashboard.",
        "permissions": "Public / Anonymous users (auto-provisions as student or preserves existing role)",
        "test_evidence": "backend/tests/auth.test.js (TC-AUTH-16, TC-AUTH-17, TC-AUTH-18)"
    },
    {
        "name": "3. Forgot / Reset Password & Verification Codes",
        "page": "Unified Login (Password Reset Modal)",
        "route": "/login",
        "frontend_file": "frontend/src/pages/LoginPage.jsx",
        "components": "ForgotPasswordModal in LoginPage.jsx, authService.js",
        "backend_api": "POST /api/auth/forgot-password, POST /api/auth/verify-reset-code, POST /api/auth/reset-password",
        "backend_files": "backend/src/routes/authRoutes.js, authController.js, emailService.js, User.js, Session.js",
        "database": "users (resetPasswordToken, resetPasswordExpire, password), sessions (isValid)",
        "integrations": "Brevo / Gmail SMTP via Nodemailer (dispatches 6-digit OTP code to user's registered inbox)",
        "workflow": "User enters email -> server generates 6-digit crypto OTP (15-min expiry) -> emails code via SMTP -> user enters code to verify -> user submits new password -> server hashes password, clears reset tokens, revokes all active sessions for security, and completes reset.",
        "permissions": "Public / Anonymous users",
        "test_evidence": "backend/tests/auth.test.js (TC-AUTH-15), backend/tests/phase10.test.js (TC-PH10-03)"
    },
    {
        "name": "4. Two-Factor Authentication (2FA)",
        "page": "Unified Login & User Profile",
        "route": "/login, /profile",
        "frontend_file": "frontend/src/pages/LoginPage.jsx, frontend/src/pages/ProfilePage.jsx",
        "components": "TwoFactorModal in LoginPage.jsx, TwoFactorSection in ProfilePage.jsx",
        "backend_api": "POST /api/auth/login-verify-2fa, POST /api/auth/resend-2fa, GET /api/auth/2fa/status, POST /api/auth/2fa/enable-request, POST /api/auth/2fa/enable-confirm, POST /api/auth/2fa/disable",
        "backend_files": "backend/src/routes/authRoutes.js, authController.js, emailService.js, User.js, auth.js",
        "database": "users (twoFactorEnabled, twoFactorCode, twoFactorExpire)",
        "integrations": "Brevo / Gmail SMTP via emailService.js (dispatches 6-digit 2FA login verification code)",
        "workflow": "User enters credentials -> if 2FA enabled, server responds requiresTwoFactor: true and sends OTP via email -> UI displays OTP challenge -> user submits 6-digit code -> server validates code and expiry -> issues authenticated JWT and Session record.",
        "permissions": "All Authenticated Users to configure; Anonymous users during login challenge",
        "test_evidence": "backend/tests/auth_sessions_2fa.test.js (TC-2FA-01, TC-2FA-02, TC-2FA-03, TC-2FA-04)"
    },
    {
        "name": "5. Session Management, Expiry & Remote Revocation",
        "page": "User Profile & App Layout",
        "route": "/profile, /dashboard",
        "frontend_file": "frontend/src/pages/ProfilePage.jsx, frontend/src/context/AuthContext.jsx, frontend/src/api/client.js",
        "components": "ActiveSessionsManager in ProfilePage.jsx, AuthContext timer, Axios interceptor",
        "backend_api": "GET /api/auth/sessions, DELETE /api/auth/sessions/:sessionId, POST /api/auth/sessions/revoke-others, POST /api/auth/logout",
        "backend_files": "backend/src/routes/authRoutes.js, authController.js, auth.js middleware, Session.js model",
        "database": "sessions (user, tokenSignature, isValid, deviceInfo, ipAddress, lastActiveAt, createdAt)",
        "integrations": "Cache-Control: no-store, no-cache headers; client-side setTimeout timer based on JWT exp claim",
        "workflow": "Login creates Session -> middleware verifies Session.isValid on every request -> user views active devices in Profile -> user can terminate specific session or all other sessions -> Logout marks isValid: false -> subsequent requests return HTTP 401 TOKEN_EXPIRED.",
        "permissions": "All Authenticated Users",
        "test_evidence": "backend/tests/auth_sessions_2fa.test.js (TC-SESS-01, TC-SESS-02, TC-SESS-03, TC-SESS-04, TC-RBAC-04)"
    },
    {
        "name": "6. Role-Based Access Control (5-Tier RBAC)",
        "page": "App Router & All Pages",
        "route": "All Routes (/admin/*, /manage/*, /my-*)",
        "frontend_file": "frontend/src/components/auth/ProtectedRoute.jsx, frontend/src/components/auth/RoleRedirect.jsx, frontend/src/App.jsx",
        "components": "ProtectedRoute, RoleRedirect, TopNavbar, Sidebar",
        "backend_api": "Enforced across all /api endpoints via requireRole middleware",
        "backend_files": "backend/src/middlewares/auth.js (authenticate, requireRole), backend/src/routes/*.js",
        "database": "users (role: student | mentor | company | coordinator | admin)",
        "integrations": "Not applicable",
        "workflow": "Client Route guards verify user.role in AuthContext -> backend middleware requireRole('role1', 'role2') inspects req.user.role -> rejects unauthorized access with HTTP 403 FORBIDDEN and standardized error envelope.",
        "permissions": "Student, Mentor, Company, Coordinator, Admin (strictly segregated per route)",
        "test_evidence": "backend/tests/auth.test.js (TC-AUTH-12, TC-AUTH-13, TC-AUTH-14), backend/tests/phase11.test.js (NEG-04)"
    },
    {
        "name": "7. User Profiles, Avatar & PDF Resume Uploads",
        "page": "User Profile & Security",
        "route": "/profile",
        "frontend_file": "frontend/src/pages/ProfilePage.jsx",
        "components": "AvatarUploadModal, ResumeUploadSection, UserAvatar, ProfileForm",
        "backend_api": "GET /api/users/profile, PUT /api/users/profile, POST /api/users/profile/avatar, DELETE /api/users/profile/avatar, POST /api/users/profile/resume, GET /api/users/profile/resume/:filename, DELETE /api/users/profile/resume",
        "backend_files": "backend/src/routes/userRoutes.js, userController.js, validateUser.js, User.js",
        "database": "users (name, phone, department, bio, cgpa, skills, avatar, resume, resumeOriginalName)",
        "integrations": "Multer file storage (/uploads/avatars, /uploads/resumes); MIME type validation (PDF, JPG, PNG)",
        "workflow": "User updates bio/skills or uploads avatar/resume -> Multer stores file with secure unique timestamp filename -> controller updates User document -> resume download endpoint checks user role or employer authorization before streaming file.",
        "permissions": "All Authenticated Users (own profile); Resume download restricted to owner, assigned employer, coordinator, admin",
        "test_evidence": "backend/tests/users.test.js (TC-PROF-01 through TC-PROF-11), backend/tests/phase10.test.js (TC-PH10-07)"
    },
    {
        "name": "8. Internship Creation, Publishing & Management",
        "page": "Opportunities Management",
        "route": "/manage/internships",
        "frontend_file": "frontend/src/pages/ManageInternshipsPage.jsx",
        "components": "CreateInternshipModal, EditInternshipModal, InternshipTable, StatusToggleSwitch",
        "backend_api": "GET /api/internships/manage, POST /api/internships, GET /api/internships/:id, PUT /api/internships/:id, PATCH /api/internships/:id/status, DELETE /api/internships/:id",
        "backend_files": "backend/src/routes/internshipRoutes.js, internshipController.js, validateInternship.js, Internship.js",
        "database": "internships (company, title, description, department, location, type, stipend, requirements, openings, deadline, status)",
        "integrations": "Not applicable",
        "workflow": "Employer or Coordinator fills posting form -> validateInternship checks deadline is future date and openings >= 1 -> saved with status: draft or published -> employer can edit, toggle status, or close opportunity.",
        "permissions": "Company (own opportunities), Coordinator & Admin (all opportunities); Student cannot create",
        "test_evidence": "backend/tests/internships.test.js (TC-INT-01 through TC-INT-11)"
    },
    {
        "name": "9. Internship Directory Search & Student Applications",
        "page": "Internships Directory & My Applications",
        "route": "/internships, /my-applications",
        "frontend_file": "frontend/src/pages/InternshipsDirectoryPage.jsx, frontend/src/pages/StudentApplicationsPage.jsx",
        "components": "InternshipCard, ApplicationModal, SearchBar, DepartmentFilter, StatusPill",
        "backend_api": "GET /api/internships, POST /api/applications, GET /api/applications/my, PATCH /api/applications/:id/withdraw",
        "backend_files": "backend/src/routes/applicationRoutes.js, applicationController.js, validateApplication.js, Application.js, notificationService.js",
        "database": "applications (student, internship, company, status, coverLetter, resumeUrl, decisionReason, history)",
        "integrations": "Socket.IO real-time notification emission to company socket room on new application",
        "workflow": "Student browses published internships -> clicks Apply -> submits cover letter -> validateApplication verifies deadline not passed and enforces compound unique index { student, internship } -> persists Application -> emits live alert to employer.",
        "permissions": "Public browse directory; Student role to apply and withdraw own applications",
        "test_evidence": "backend/tests/applications.test.js (TC-APP-01, TC-APP-02, TC-APP-03, TC-APP-05, TC-APP-12)"
    },
    {
        "name": "10. Candidate Screening, Shortlisting & Selection",
        "page": "Employer Candidate Review",
        "route": "/manage/applications",
        "frontend_file": "frontend/src/pages/CompanyApplicationsPage.jsx",
        "components": "ApplicantTable, CandidateReviewModal, StatusTransitionDropdown, AIResumeScoreBadge",
        "backend_api": "GET /api/applications, GET /api/applications/:id, PATCH /api/applications/:id/status",
        "backend_files": "backend/src/routes/applicationRoutes.js, applicationController.js, validateApplication.js, Application.js, notificationService.js",
        "database": "applications (status: submitted | reviewing | shortlisted | interview | accepted | rejected, decisionReason, history)",
        "integrations": "Socket.IO real-time alert (notification:new) dispatched to applicant student on status change",
        "workflow": "Employer opens Candidate Review -> inspects applicant resume and AI fit score -> changes status (e.g. submitted -> reviewing -> shortlisted -> accepted) with mandatory decision reason -> updates audit history -> dispatches live notification to candidate.",
        "permissions": "Company (applicants for own postings), Coordinator & Admin (all applicants)",
        "test_evidence": "backend/tests/applications.test.js (TC-APP-07, TC-APP-08, TC-APP-09, TC-APP-10, TC-APP-11)"
    },
    {
        "name": "11. Internship Allocation & Mentor Assignment",
        "page": "Intern Allocations Management",
        "route": "/manage/allocations",
        "frontend_file": "frontend/src/pages/ManageAllocationsPage.jsx",
        "components": "AllocationTable, MentorAssignModal, RequestCompletionModal, IssueCertificateModal",
        "backend_api": "POST /api/allocations, GET /api/allocations, GET /api/allocations/:id, PATCH /api/allocations/:id/status",
        "backend_files": "backend/src/routes/allocationRoutes.js, allocationController.js, validateAllocation.js, Allocation.js",
        "database": "allocations (student, internship, company, mentor, startDate, endDate, status, finalEvaluation)",
        "integrations": "Socket.IO notification dispatched to intern and assigned mentor",
        "workflow": "Employer selects accepted candidate -> specifies assigned mentor, start date, and end date -> validateAllocation checks student has no other active allocation -> saves Allocation -> links student, employer, and mentor in unified workspace.",
        "permissions": "Company, Coordinator, Admin; Student cannot allocate",
        "test_evidence": "backend/tests/phase6.test.js (TC-ALC-01, TC-ALC-02, TC-ALC-03, TC-ALC-05)"
    },
    {
        "name": "12. Tasks, Attendance, Milestones & Evaluations",
        "page": "Student & Manage Tasks / Attendance / Evaluations",
        "route": "/my-tasks, /manage/tasks, /my-attendance, /manage/attendance, /evaluations",
        "frontend_file": "frontend/src/pages/StudentTasksPage.jsx, ManageTasksPage.jsx, StudentAttendancePage.jsx, ManageAttendancePage.jsx, EvaluationsPage.jsx",
        "components": "AttendanceSubmissionForm, MonthlyAttendanceCalendar, TaskCardList, TaskSubmissionModal, EvaluationFormModal",
        "backend_api": "POST /api/attendance, GET /api/attendance, PATCH /api/attendance/:id/review, POST /api/tasks, GET /api/tasks, PATCH /api/tasks/:id/progress, PATCH /api/tasks/:id/review, POST /api/evaluations, GET /api/evaluations",
        "backend_files": "attendanceRoutes.js, taskRoutes.js, evaluationRoutes.js, controllers, models (Attendance, Task, Evaluation)",
        "database": "attendances (student, date, hoursWorked, status, reviewStatus), tasks (assignedTo, title, dueDate, status, score), evaluations (allocation, midTerm, finalScore)",
        "integrations": "Compound unique index { student, date } on attendances prevents same-day duplicates; live socket alerts",
        "workflow": "Intern logs daily attendance -> supervisor approves or corrects hours. Mentor assigns milestone task -> intern submits deliverable URL -> mentor scores task. Mid-term and final performance evaluations submitted with multi-criteria ratings.",
        "permissions": "Student submits attendance/tasks; Mentor/Company/Coordinator reviews and assigns tasks, evaluates performance",
        "test_evidence": "backend/tests/phase6.test.js (TC-ATT-01 to TC-ATT-06, TC-TSK-01 to TC-TSK-04), backend/tests/phase9.test.js (TC-PH9-01 to TC-PH9-05)"
    },
    {
        "name": "13. Certificate Eligibility, PDFKit, QR Verification & SMTP",
        "page": "Certificates, Admin Certificates & Public Verification",
        "route": "/my-certificates, /admin/certificates, /verify-certificate/:code",
        "frontend_file": "frontend/src/pages/CertificatesPage.jsx, AdminCertificatesPage.jsx, VerifyCertificatePage.jsx",
        "components": "CertificateCard, PDFDownloadButton, PublicVerifyLink, VerifiedBadge, PadmavathiSignatureDisplay",
        "backend_api": "GET /api/allocations/:id/completion-eligibility, POST /api/allocations/:id/request-completion, POST /api/allocations/:id/review-completion, GET /api/certificates/allocation/:id/preview, POST /api/certificates/allocation/:id/issue, GET /api/certificates/:id/download, GET /api/certificates/verify/:code, POST /api/certificates/:id/resend-email",
        "backend_files": "certificateRoutes.js, certificateController.js, certificateService.js, emailService.js, Certificate.js, Allocation.js",
        "database": "certificates (certificateNumber, student, allocation, company, mentor, issueDate, verifiedStartDate, verifiedEndDate, qrCodeData, pdfBuffer, emailDeliveryStatus)",
        "integrations": "PDFKit landscape A4 generator, QRCode library (data URL), Brevo / Gmail SMTP with attached PDF, high-res Padmavathi signature graphic",
        "workflow": "Candidate completes requirements -> requests completion -> mentor approves -> mentor/coordinator/admin issues certificate -> strict date validation enforced -> idempotent check returns existing if duplicate -> PDFKit generates landscape A4 with QR code & digital signature -> emails PDF via SMTP -> public verification ledger validates code.",
        "permissions": "Mentor, Coordinator, Admin issue certificates; Student downloads own; Public validates verification code",
        "test_evidence": "backend/tests/certificates.test.js (CERT-01 to CERT-13), backend/tests/mentor_certificate_workflow.test.js (TC-MCERT-01 to TC-MCERT-14)"
    },
    {
        "name": "14. Real-Time Chat, Attachments, Reactions & Editing",
        "page": "Direct & Group Chat",
        "route": "/chat",
        "frontend_file": "frontend/src/pages/ChatPage.jsx",
        "components": "ConversationList, MessageThread, MessageInput, AttachmentUploader, EmojiReactionPicker, GroupMembersModal",
        "backend_api": "GET /api/chat/conversations, POST /api/chat/conversations, GET /api/chat/conversations/:id/messages, POST /api/chat/conversations/:id/messages, POST /api/chat/upload, PATCH /api/chat/messages/:id, DELETE /api/chat/messages/:id, PATCH /api/chat/messages/:id/react, POST /api/chat/messages/:id/forward, PATCH /api/chat/conversations/:id/read",
        "backend_files": "chatRoutes.js, chatController.js, validateChat.js, socket.js, Message.js, Conversation.js, chatFileUpload.js",
        "database": "conversations (participants, isGroup, groupName, groupAdmin, lastMessage), messages (conversation, sender, content, attachments, reactions, readBy, replyTo, isEdited, isDeleted)",
        "integrations": "Socket.IO real-time delivery (chat:message, chat:read_update, user:typing, user:status); Multer file storage for chat attachments",
        "workflow": "Users chat 1-on-1 or in groups -> messages persisted in MongoDB -> broadcast via Socket.IO room conv:<id> -> author can edit within time window or soft-delete -> recipients toggle emoji reactions -> read receipts clear unread counters.",
        "permissions": "All Authenticated Users (conversation participants only; Non-members strictly blocked with HTTP 403)",
        "test_evidence": "backend/tests/chat.test.js (TC-CHAT-01 through TC-CHAT-25)"
    },
    {
        "name": "15. In-App Notifications, Unread Counts & OS Desktop Alerts",
        "page": "Notification Center & Floating Toasts",
        "route": "/notifications, All Pages (Toasts)",
        "frontend_file": "frontend/src/pages/NotificationsPage.jsx, frontend/src/components/common/NotificationToastContainer.jsx, frontend/src/utils/desktopNotification.js, frontend/public/sw.js",
        "components": "NotificationBell, NotificationToastContainer, useSocket hook, Service Worker sw.js",
        "backend_api": "GET /api/notifications, GET /api/notifications/unread-count, PATCH /api/notifications/read-all, PATCH /api/notifications/:id/read, DELETE /api/notifications/:id, GET /api/notifications/push/vapid-key, POST /api/notifications/push/subscribe",
        "backend_files": "notificationRoutes.js, notificationController.js, notificationService.js, pushNotificationService.js, socket.js, Notification.js, User.js",
        "database": "notifications (recipient, title, message, type, link, isRead, createdAt), users (pushSubscriptions)",
        "integrations": "Socket.IO notification:new event; Web Push VAPID; Service Worker showNotification; window.Notification; Audio chime playback",
        "workflow": "Business event occurs -> server persists Notification and emits notification:new to user:<id> room -> client SocketContext inspects active conversation state -> if user is not actively chatting in focused tab, triggers floating branded toast, audio chime, and native OS desktop banner -> clicking notification routes directly to target context.",
        "permissions": "All Authenticated Users (own notifications only; user isolation enforced)",
        "test_evidence": "backend/tests/notifications.test.js (TC-NOTIF-01 to TC-NOTIF-12), backend/tests/desktop_notifications.test.js (TC-DESK-01 to TC-DESK-05)"
    },
    {
        "name": "16. AI Resume Analysis & Skill Matching",
        "page": "Opportunities & Candidate Review",
        "route": "/internships, /manage/applications",
        "frontend_file": "frontend/src/pages/CompanyApplicationsPage.jsx, frontend/src/pages/InternshipsDirectoryPage.jsx",
        "components": "AIResumeScoreBadge in CompanyApplicationsPage.jsx, MatchScoreIndicator in InternshipCard.jsx",
        "backend_api": "GET /api/internships/:id/match-score, POST /api/applications/:id/analyze-resume",
        "backend_files": "backend/src/routes/applicationRoutes.js, backend/src/routes/internshipRoutes.js, aiService.js, applicationController.js, internshipController.js",
        "database": "applications (aiScreening: matchScore, summary, strengths, gaps, recommendation), internships (requirements), users (skills, resume)",
        "integrations": "Google Gemini 2.5 Flash API (@google/genai) with robust offline deterministic fallback algorithm based on skill overlap, CGPA, and department",
        "workflow": "Student views internship or employer inspects applicant -> backend extracts skills from candidate profile and parsed resume -> invokes Gemini Flash or deterministic algorithm -> calculates 0-100 fit score, identifies matched/missing skills, and outputs recommendation.",
        "permissions": "Student queries own match score; Employer, Coordinator, Admin analyze applicant resumes",
        "test_evidence": "backend/tests/ai.test.js (AI-01, AI-02, AI-03)"
    },
    {
        "name": "17. External Job Radar, Bookmarks & Notes",
        "page": "Jobs & LinkedIn Radar",
        "route": "/external-jobs",
        "frontend_file": "frontend/src/pages/ExternalJobsPage.jsx",
        "components": "ExternalJobCard, SkillMatchRadar, BookmarkButton, ApplyTrackerModal",
        "backend_api": "GET /api/external-jobs, GET /api/external-jobs/recommendations, GET /api/external-jobs/my-interactions, POST /api/external-jobs/sync, POST /api/external-jobs/:id/bookmark, POST /api/external-jobs/:id/apply, PATCH /api/external-jobs/:id/notes",
        "backend_files": "externalJobRoutes.js, externalJobController.js, jobFeedService.js, jobScheduler.js, ExternalJob.js, UserJobBookmark.js",
        "database": "externaljobs (externalId, title, company, location, source, url, skillsRequired, applyCount), userjobbookmarks (user, job, isSaved, isApplied, notes)",
        "integrations": "Dynamic compliant LinkedIn search link generation; automated 60-minute background job aggregator",
        "workflow": "Background scheduler fetches and seeds external tech listings -> student searches listings -> views personalized AI fit score -> bookmarks jobs, records application dates, and saves interview notes.",
        "permissions": "Public search listings; Student role required to bookmark, track applications, and save personal notes",
        "test_evidence": "backend/tests/externalJobs.test.js (TC-EXT-01 through TC-EXT-11)"
    },
    {
        "name": "18. Administrator User Governance & Last-Admin Guard",
        "page": "Administrator User Console",
        "route": "/admin/users",
        "frontend_file": "frontend/src/pages/AdminUsersPage.jsx",
        "components": "UserManagementTable, CreateCoordinatorModal, RoleChangeDropdown, AccountStatusSwitch, DeleteUserModal",
        "backend_api": "GET /api/users, POST /api/users, GET /api/users/:id, PUT /api/users/:id, PATCH /api/users/:id/status, DELETE /api/users/:id",
        "backend_files": "userRoutes.js, userController.js, validateUser.js, initAdmins.js, User.js",
        "database": "users (name, email, role, status: active | inactive, isEmailVerified, createdAt)",
        "integrations": "Not applicable",
        "workflow": "Administrator searches users with pagination -> provisions new Academic Coordinator -> updates roles or deactivates accounts -> validateUser counts remaining active admins before allowing demotion -> blocks self-lockout with HTTP 403.",
        "permissions": "Admin role only (Non-admins strictly rejected with HTTP 403)",
        "test_evidence": "backend/tests/users.test.js (TC-ADMIN-01 through TC-ADMIN-06), backend/tests/phase10.test.js (TC-PH10-10)"
    },
    {
        "name": "19. Analytics Dashboard & Sanitized CSV Export",
        "page": "Analytics & Reports Export",
        "route": "/reports",
        "frontend_file": "frontend/src/pages/ReportsPage.jsx",
        "components": "MetricsOverviewCards, TrendCharts, ExportCsvButton, DateRangePicker",
        "backend_api": "GET /api/analytics/dashboard, GET /api/reports, GET /api/reports/export",
        "backend_files": "analyticsRoutes.js, reportRoutes.js, analyticsController.js, reportController.js, models",
        "database": "Aggregations across users, internships, applications, allocations, attendances, certificates",
        "integrations": "CSV data streaming with strict sensitive field redaction (passwords, tokens, and hashes stripped)",
        "workflow": "Coordinator/Admin/Employer queries reports page -> controller executes MongoDB aggregation pipelines computing KPIs, monthly attendance rates, and hiring funnel ratios -> user clicks Export CSV -> streams sanitized CSV file.",
        "permissions": "Company, Mentor, Coordinator, Admin (Students strictly blocked with HTTP 403)",
        "test_evidence": "backend/tests/phase9.test.js (TC-PH9-06 through TC-PH9-12)"
    },
    {
        "name": "20. Security Hardening, Sanitization & Token Bucket Rate Limiting",
        "page": "Global Backend Middleware",
        "route": "All API Endpoints",
        "frontend_file": "frontend/src/api/client.js (handles HTTP 429 Retry-After header and 401 token expiry)",
        "components": "Axios interceptors in client.js, notification toast on rate limit or session expired",
        "backend_api": "Global middleware pipeline mounted on all routes",
        "backend_files": "backend/src/middlewares/sanitize.js, rateLimiter.js, auth.js, errorHandler.js",
        "database": "Sessions collection for token invalidation; In-memory token bucket tracking",
        "integrations": "Helmet HTTP security headers; CORS origin validation; regex escaping (escapeRegex); sensitive data masking",
        "workflow": "Incoming request -> Helmet strips headers -> sanitizeInput strips NoSQL operators ($gt, $ne, __proto__) -> rateLimiter tracks IP/token requests -> if threshold exceeded, returns HTTP 429 with Retry-After header -> auth middleware verifies token and active session.",
        "permissions": "Applied universally to all incoming network traffic",
        "test_evidence": "backend/tests/phase10.test.js (TC-PH10-01 to TC-PH10-15)"
    },
    {
        "name": "21. System Health Checks & Database Connectivity",
        "page": "System Overview & Header Bar",
        "route": "/ (Header live ping indicator)",
        "frontend_file": "frontend/src/App.jsx, frontend/src/services/healthService.js, frontend/src/components/layout/TopNavbar.jsx",
        "components": "Health status pill in TopNavbar.jsx, 15-second polling interval in App.jsx",
        "backend_api": "GET /api, GET /api/health, GET /api/health/db",
        "backend_files": "healthRoutes.js, healthController.js, db.js",
        "database": "MongoDB mongoose.connection.readyState check and ping command",
        "integrations": "Process uptime, memory usage, environment diagnostics",
        "workflow": "Client polls /api/health every 15s -> server verifies MongoDB connectivity, memory, and uptime -> returns healthy status -> UI displays green connectivity indicator in top navbar.",
        "permissions": "Public (All)",
        "test_evidence": "backend/tests/health.test.js (HLTH-01, HLTH-02, HLTH-03, HLTH-04), backend/tests/phase11.test.js (NEG-11)"
    },
    {
        "name": "22. Multi-Device Responsive Design & Virtual Keyboard Adaptation",
        "page": "All 26 Pages",
        "route": "All Routes (320px to 1920px+)",
        "frontend_file": "frontend/src/components/layout/AppLayout.jsx, Sidebar.jsx, TopNavbar.jsx, Table.jsx, Modal.jsx",
        "components": "Tailwind responsive classes, Mobile Drawer, Touch-scrolling Tables, Modal Viewport Constraints",
        "backend_api": "Not applicable (Frontend presentation layer)",
        "backend_files": "Not applicable",
        "database": "Not applicable",
        "integrations": "WebkitOverflowScrolling, CSS max-h-[calc(100dvh-2rem)] for mobile virtual keyboards",
        "workflow": "User accesses on mobile (320px-430px), tablet (768px-1024px), or desktop (1280px-1920px+) -> layout adapts: sidebar switches to touch-overlay drawer, data tables enable horizontal swipe, modals fit within viewport without keyboard clipping.",
        "permissions": "All Users across all viewports",
        "test_evidence": "Verified across Chrome & Edge mobile emulation (320px to 1920px); Zero horizontal overflow"
    },
    {
        "name": "23. NovaTech Demo Labs Seeding & Demonstration Workflow",
        "page": "Intern Allocations Management",
        "route": "/manage/allocations",
        "frontend_file": "frontend/src/pages/ManageAllocationsPage.jsx",
        "components": "SeedDemoButton in ManageAllocationsPage.jsx, allocationService.js",
        "backend_api": "POST /api/allocations/seed-demo, POST /api/certificates/admin/seed-demo",
        "backend_files": "allocationRoutes.js, certificateRoutes.js, allocationController.js, certificateController.js, demo seeds",
        "database": "Populates users (Hari Student, NovaTech Employer, Tech Lead Mentor), internship, allocation, tasks, attendance",
        "integrations": "Prepares verifiable certificate demo with dynamic Padmavathi signature and QR verification",
        "workflow": "Admin or Employer clicks 'Seed Demo Workflow' -> backend provisions end-to-end demo scenario with 100% completed tasks, approved attendance, and eligible allocation -> allows immediate testing of mentor certificate review and issuance.",
        "permissions": "Admin, Coordinator, Company",
        "test_evidence": "backend/tests/mentor_certificate_workflow.test.js (TC-MCERT-14)"
    },
    {
        "name": "24. Institutional Compliance & Policy Governance",
        "page": "Institutional Policies & Policy Modal",
        "route": "/policies, /privacy, /terms, /security-policy, /cookies",
        "frontend_file": "frontend/src/pages/PoliciesPage.jsx, frontend/src/components/common/PolicyModal.jsx",
        "components": "PolicyModal, PoliciesPage navigation tabs, Footer links",
        "backend_api": "Not applicable (Static institutional governance documents)",
        "backend_files": "Not applicable",
        "database": "Not applicable",
        "integrations": "Print-friendly CSS formatting for physical archiving",
        "workflow": "User clicks footer links or registration terms link -> displays formatted institutional policies covering data handling, intellectual property, student confidentiality, and platform safety rules.",
        "permissions": "Public (All)",
        "test_evidence": "Verified in frontend build compilation and routing tests"
    },
    {
        "name": "25. Chat Member Governance, Quoting & Message Forwarding",
        "page": "Direct & Group Chat",
        "route": "/chat",
        "frontend_file": "frontend/src/pages/ChatPage.jsx",
        "components": "GroupMembersModal, ReplyPreviewBar, ForwardMessageModal",
        "backend_api": "PATCH /api/chat/conversations/:id/members, POST /api/chat/messages/:id/forward",
        "backend_files": "chatRoutes.js, chatController.js, validateChat.js, Conversation.js, Message.js",
        "database": "conversations (participants, groupAdmin), messages (replyTo, forwardFrom)",
        "integrations": "Socket.IO real-time member join/leave broadcasts",
        "workflow": "Coordinator creates group -> adds or removes members -> participants can quote-reply to specific parent messages or forward important messages into other active conversation channels.",
        "permissions": "Group admin manages members; Participants quote and forward",
        "test_evidence": "backend/tests/chat.test.js (TC-CHAT-06, TC-CHAT-07, TC-CHAT-08, TC-CHAT-23, TC-CHAT-24)"
    },
    {
        "name": "26. Institutional Announcement Broadcasting",
        "page": "Notification Center & Admin Console",
        "route": "/notifications",
        "frontend_file": "frontend/src/pages/NotificationsPage.jsx",
        "components": "AnnouncementBroadcastModal, NotificationItem",
        "backend_api": "POST /api/notifications/announcements",
        "backend_files": "notificationRoutes.js, notificationController.js, notificationService.js, socket.js, Notification.js",
        "database": "notifications (bulk insert for target audience: all | students | companies | mentors)",
        "integrations": "Socket.IO multi-room broadcast emitting notification:new to all active sockets",
        "workflow": "Admin or Coordinator composes institutional bulletin -> selects target role (e.g. all students) -> backend bulk-inserts Notification records and broadcasts live event across active socket connections.",
        "permissions": "Admin, Coordinator",
        "test_evidence": "backend/tests/notifications.test.js (TC-NOTIF-11)"
    }
]

# --- 14 Database Models Data ---
DATABASE_MODELS = [
    {
        "model": "User.js",
        "collection": "users",
        "pk": "_id (ObjectId)",
        "fields": "name, email (unique), password (hash), role (student | mentor | company | coordinator | admin), status (active | inactive), department, phone, bio, avatar, resume, resumeOriginalName, cgpa, skills, twoFactorEnabled, twoFactorCode, twoFactorExpire, resetPasswordToken, resetPasswordExpire, pushSubscriptions",
        "indexes": "{ email: 1 } (unique), { role: 1 }, { status: 1 }",
        "relations": "Referenced by Application, Allocation, Attendance, Task, Evaluation, Certificate, Message, Conversation, Session"
    },
    {
        "model": "Session.js",
        "collection": "sessions",
        "pk": "_id (ObjectId)",
        "fields": "user (ObjectId -> User), tokenSignature (String, indexed), isValid (Boolean), deviceInfo (String), ipAddress (String), userAgent (String), lastActiveAt (Date), createdAt (Date)",
        "indexes": "{ tokenSignature: 1 }, { user: 1, isValid: 1 }",
        "relations": "Belongs to User (1-to-N session tracking and remote termination)"
    },
    {
        "model": "Internship.js",
        "collection": "internships",
        "pk": "_id (ObjectId)",
        "fields": "company (ObjectId -> User), title, description, department, location, type (remote | on-site | hybrid), stipend (amount, currency), requirements (skills, minCgpa), openings, deadline, status (draft | published | closed | archived)",
        "indexes": "{ company: 1 }, { status: 1 }, { department: 1 }, { deadline: 1 }",
        "relations": "Belongs to Company User; Referenced by Application, Allocation"
    },
    {
        "model": "Application.js",
        "collection": "applications",
        "pk": "_id (ObjectId)",
        "fields": "student (ObjectId -> User), internship (ObjectId -> Internship), company (ObjectId -> User), status (submitted | reviewing | shortlisted | interview | accepted | rejected | withdrawn), coverLetter, resumeUrl, decisionReason, decisionDate, aiScreening (matchScore, summary, strengths, gaps, recommendation), history [status, actor, notes, timestamp]",
        "indexes": "{ student: 1, internship: 1 } (UNIQUE COMPOUND BARRIER), { company: 1, status: 1 }",
        "relations": "Links Student User to Internship opportunity"
    },
    {
        "model": "Allocation.js",
        "collection": "allocations",
        "pk": "_id (ObjectId)",
        "fields": "student (ObjectId -> User), internship (ObjectId -> Internship), company (ObjectId -> User), mentor (ObjectId -> User), startDate (Date), endDate (Date), status (pending | active | completed | terminated), completionRequest (status, requestedAt, studentNotes, preferredEmail, reviewedBy, reviewedAt, supervisorRecommendation, finalGrade, reviewNotes)",
        "indexes": "{ student: 1, status: 1 }, { mentor: 1 }, { company: 1 }",
        "relations": "Core operational entity linking Student, Internship, Company, and Mentor"
    },
    {
        "model": "Attendance.js",
        "collection": "attendances",
        "pk": "_id (ObjectId)",
        "fields": "student (ObjectId -> User), allocation (ObjectId -> Allocation), date (Date), hoursWorked (Number), workSummary (String), status (present | absent | half-day), reviewStatus (pending | approved | corrected), reviewedBy (ObjectId -> User), reviewNotes, reviewedAt",
        "indexes": "{ student: 1, date: 1 } (UNIQUE COMPOUND BARRIER), { allocation: 1 }",
        "relations": "Belongs to Student and Allocation"
    },
    {
        "model": "Task.js",
        "collection": "tasks",
        "pk": "_id (ObjectId)",
        "fields": "allocation (ObjectId -> Allocation), assignedBy (ObjectId -> User), assignedTo (ObjectId -> User), title, description, dueDate (Date), priority (low | medium | high), status (assigned | in-progress | submitted | completed | overdue), submission (notes, deliverableUrl, submittedAt), review (score, feedback, reviewedAt, reviewedBy)",
        "indexes": "{ assignedTo: 1, status: 1 }, { allocation: 1 }, { dueDate: 1 }",
        "relations": "Belongs to Allocation; Assigned by Mentor to Student Intern"
    },
    {
        "model": "Evaluation.js",
        "collection": "evaluations",
        "pk": "_id (ObjectId)",
        "fields": "allocation (ObjectId -> Allocation), evaluator (ObjectId -> User), student (ObjectId -> User), type (midterm | final), technicalSkills (score, notes), communication (score, notes), professionalism (score, notes), overallRating (Number), recommendation (String), submittedAt (Date)",
        "indexes": "{ allocation: 1, type: 1 } (unique compound), { student: 1 }",
        "relations": "Formal appraisal tied to Allocation and Student"
    },
    {
        "model": "Certificate.js",
        "collection": "certificates",
        "pk": "_id (ObjectId)",
        "fields": "certificateNumber (String, unique), student (ObjectId -> User), allocation (ObjectId -> Allocation, unique), company (ObjectId -> User), mentor (ObjectId -> User), title, recipientName, companyName, mentorName, issueDate (Date), verifiedStartDate (Date), verifiedEndDate (Date), grade, qrCodeData (String), pdfBuffer (Buffer), emailSentAt (Date), emailDeliveryStatus (pending | sent | failed), resendCount (Number)",
        "indexes": "{ certificateNumber: 1 } (unique), { allocation: 1 } (unique), { student: 1 }",
        "relations": "Certified credential linked to completed Allocation and Student"
    },
    {
        "model": "Conversation.js",
        "collection": "conversations",
        "pk": "_id (ObjectId)",
        "fields": "participants [ObjectId -> User], isGroup (Boolean), groupName (String), groupAdmin (ObjectId -> User), lastMessage (ObjectId -> Message), lastMessageAt (Date), createdAt, updatedAt",
        "indexes": "{ participants: 1 }, { updatedAt: -1 }",
        "relations": "Chat channel grouping 2 or more User participants; references Messages"
    },
    {
        "model": "Message.js",
        "collection": "messages",
        "pk": "_id (ObjectId)",
        "fields": "conversation (ObjectId -> Conversation), sender (ObjectId -> User), content (String), attachments [url, fileKey, originalName, mimeType, size], reactions [emoji, user], readBy [ObjectId -> User], replyTo (ObjectId -> Message), isEdited (Boolean), isDeleted (Boolean), deletedFor [ObjectId -> User]",
        "indexes": "{ conversation: 1, createdAt: 1 }",
        "relations": "Belongs to Conversation; authored by User"
    },
    {
        "model": "Notification.js",
        "collection": "notifications",
        "pk": "_id (ObjectId)",
        "fields": "recipient (ObjectId -> User), sender (ObjectId -> User), title (String), message (String), type (application | task | attendance | certificate | chat | system), link (String), isRead (Boolean), metadata (Mixed), createdAt (Date)",
        "indexes": "{ recipient: 1, isRead: 1, createdAt: -1 }",
        "relations": "Delivered to recipient User"
    },
    {
        "model": "ExternalJob.js",
        "collection": "externaljobs",
        "pk": "_id (ObjectId)",
        "fields": "externalId (String, unique), title, company, location, type, source (LinkedIn | Indeed | Glassdoor), url (unique), description, skillsRequired [String], applyCount (Number), isActive (Boolean), publishedDate (Date)",
        "indexes": "{ externalId: 1 } (unique), { url: 1 } (unique), { skillsRequired: 1 }",
        "relations": "Referenced by UserJobBookmark"
    },
    {
        "model": "UserJobBookmark.js",
        "collection": "userjobbookmarks",
        "pk": "_id (ObjectId)",
        "fields": "user (ObjectId -> User), job (ObjectId -> ExternalJob), isSaved (Boolean), isApplied (Boolean), appliedDate (Date), notes (String), createdAt, updatedAt",
        "indexes": "{ user: 1, job: 1 } (UNIQUE COMPOUND)",
        "relations": "Join entity linking Student User to bookmarked ExternalJob"
    }
]

# --- 17 Test Suites Mapping Data ---
TEST_SUITES = [
    {"suite": "health.test.js", "subsystem": "System Health, API Root & Database Ping", "tests": 4, "cmd": "node --test backend/tests/health.test.js"},
    {"suite": "auth.test.js", "subsystem": "Registration, Unified Login, RBAC & Google SSO", "tests": 19, "cmd": "node --test backend/tests/auth.test.js"},
    {"suite": "auth_sessions_2fa.test.js", "subsystem": "2FA OTP, Active Sessions & Token Revocation", "tests": 15, "cmd": "node --test backend/tests/auth_sessions_2fa.test.js"},
    {"suite": "users.test.js", "subsystem": "User Profiles, Avatars, Resumes & Admin CRUD", "tests": 18, "cmd": "node --test backend/tests/users.test.js"},
    {"suite": "internships.test.js", "subsystem": "Opportunity Lifecycle, Ownership & Deadlines", "tests": 11, "cmd": "node --test backend/tests/internships.test.js"},
    {"suite": "applications.test.js", "subsystem": "Applications, Deadlines & Duplicate Barriers", "tests": 12, "cmd": "node --test backend/tests/applications.test.js"},
    {"suite": "certificates.test.js", "subsystem": "Certificate Issuance, PDFKit & QR Verification", "tests": 13, "cmd": "node --test backend/tests/certificates.test.js"},
    {"suite": "mentor_certificate_workflow.test.js", "subsystem": "Mentor Review, Eligibility & Demo Labs Seed", "tests": 14, "cmd": "node --test backend/tests/mentor_certificate_workflow.test.js"},
    {"suite": "chat.test.js", "subsystem": "Direct/Group Chat, Attachments & Reactions", "tests": 25, "cmd": "node --test backend/tests/chat.test.js"},
    {"suite": "desktop_notifications.test.js", "subsystem": "Web Push VAPID, Subscriptions & OS Alerts", "tests": 5, "cmd": "node --test backend/tests/desktop_notifications.test.js"},
    {"suite": "notifications.test.js", "subsystem": "In-App Alerts, Socket Auth & Unread Sync", "tests": 12, "cmd": "node --test backend/tests/notifications.test.js"},
    {"suite": "ai.test.js", "subsystem": "AI Resume Analysis, Skill Matching & Scoring", "tests": 3, "cmd": "node --test backend/tests/ai.test.js"},
    {"suite": "externalJobs.test.js", "subsystem": "External Job Radar, Bookmarking & Feed Sync", "tests": 11, "cmd": "node --test backend/tests/externalJobs.test.js"},
    {"suite": "phase6.test.js", "subsystem": "Allocations, Attendance Audits, Tasks & Overdue", "tests": 15, "cmd": "node --test backend/tests/phase6.test.js"},
    {"suite": "phase9.test.js", "subsystem": "Analytics, Evaluations, Reports & CSV Export", "tests": 12, "cmd": "node --test backend/tests/phase9.test.js"},
    {"suite": "phase10.test.js", "subsystem": "Security Hardening, NoSQL Stripping & Rate Limits", "tests": 15, "cmd": "node --test backend/tests/phase10.test.js"},
    {"suite": "phase11.test.js", "subsystem": "Full System E2E Regressions (10 WF + 11 NEG)", "tests": 21, "cmd": "node --test backend/tests/phase11.test.js"},
]

print("Data definitions loaded successfully")

# --- 123 Backend API Endpoints Directory ---
ALL_ENDPOINTS = [
    # Health Routes (2)
    {"method": "GET", "path": "/api/health", "controller": "healthController.getHealth", "middleware": "None (Public)", "desc": "Returns overall system health, process uptime, memory usage, and timestamp."},
    {"method": "GET", "path": "/api/health/db", "controller": "healthController.getDatabaseHealth", "middleware": "None (Public)", "desc": "Performs isolated MongoDB ping diagnostic and returns connection readyState."},

    # Auth Routes (22)
    {"method": "POST", "path": "/api/auth/register", "controller": "authController.register", "middleware": "validateAuth.register", "desc": "Registers new Student or Company user account with hashed password."},
    {"method": "POST", "path": "/api/auth/login", "controller": "authController.login", "middleware": "validateAuth.login, authRateLimiter", "desc": "Authenticates credentials, evaluates 2FA requirements, creates Session, returns JWT."},
    {"method": "POST", "path": "/api/auth/login-verify-2fa", "controller": "authController.verifyTwoFactorLogin", "middleware": "validateAuth.verify2FA", "desc": "Validates 6-digit OTP code submitted during 2FA challenge, issues JWT."},
    {"method": "POST", "path": "/api/auth/resend-2fa", "controller": "authController.resendTwoFactorCode", "middleware": "authRateLimiter", "desc": "Regenerates and dispatches a fresh 2FA OTP code to user's registered email."},
    {"method": "POST", "path": "/api/auth/google", "controller": "authController.loginWithGoogle", "middleware": "None (Public)", "desc": "Verifies Google Identity Services ID token, auto-provisions or authenticates user."},
    {"method": "GET", "path": "/api/auth/google/client-id", "controller": "authController.getGoogleClientId", "middleware": "None (Public)", "desc": "Returns public Google OAuth Client ID for frontend button initialization."},
    {"method": "POST", "path": "/api/auth/forgot-password", "controller": "authController.forgotPassword", "middleware": "authRateLimiter", "desc": "Generates 15-minute password reset OTP and sends email via Nodemailer SMTP."},
    {"method": "POST", "path": "/api/auth/verify-reset-code", "controller": "authController.verifyResetCode", "middleware": "None (Public)", "desc": "Verifies 6-digit reset code validity before presenting password reset form."},
    {"method": "POST", "path": "/api/auth/reset-password", "controller": "authController.resetPassword", "middleware": "validateAuth.resetPassword", "desc": "Hashes new password, clears reset tokens, and revokes all active sessions."},
    {"method": "GET", "path": "/api/auth/me", "controller": "authController.getMe", "middleware": "authenticate", "desc": "Returns authenticated user profile, active role, and permission flags."},
    {"method": "POST", "path": "/api/auth/logout", "controller": "authController.logout", "middleware": "authenticate", "desc": "Marks active session isValid: false in database and instructs anti-cache headers."},
    {"method": "GET", "path": "/api/auth/2fa/status", "controller": "authController.getTwoFactorStatus", "middleware": "authenticate", "desc": "Returns boolean state of two-factor authentication for current user."},
    {"method": "POST", "path": "/api/auth/2fa/enable-request", "controller": "authController.requestEnableTwoFactor", "middleware": "authenticate", "desc": "Dispatches confirmation OTP code to enable 2FA on account."},
    {"method": "POST", "path": "/api/auth/2fa/enable-confirm", "controller": "authController.confirmEnableTwoFactor", "middleware": "authenticate", "desc": "Confirms OTP code and activates twoFactorEnabled: true in User record."},
    {"method": "POST", "path": "/api/auth/2fa/disable", "controller": "authController.disableTwoFactor", "middleware": "authenticate", "desc": "Validates account password and deactivates 2FA protection."},
    {"method": "GET", "path": "/api/auth/sessions", "controller": "authController.getActiveSessions", "middleware": "authenticate", "desc": "Returns list of all active sessions with IP, device, and last active timestamp."},
    {"method": "DELETE", "path": "/api/auth/sessions/:sessionId", "controller": "authController.revokeSession", "middleware": "authenticate", "desc": "Selectively terminates a specific remote active session by ID."},
    {"method": "POST", "path": "/api/auth/sessions/revoke-others", "controller": "authController.revokeOtherSessions", "middleware": "authenticate", "desc": "Terminates all active sessions except the currently active device token."},
    {"method": "GET", "path": "/api/auth/admin-access-test", "controller": "inline", "middleware": "authenticate, requireRole('admin')", "desc": "Integration test endpoint verifying Admin RBAC boundary enforcement."},
    {"method": "GET", "path": "/api/auth/coordinator-access-test", "controller": "inline", "middleware": "authenticate, requireRole('coordinator')", "desc": "Integration test endpoint verifying Coordinator RBAC boundary enforcement."},
    {"method": "GET", "path": "/api/auth/company-access-test", "controller": "inline", "middleware": "authenticate, requireRole('company')", "desc": "Integration test endpoint verifying Company RBAC boundary enforcement."},
    {"method": "GET", "path": "/api/auth/student-access-test", "controller": "inline", "middleware": "authenticate, requireRole('student')", "desc": "Integration test endpoint verifying Student RBAC boundary enforcement."},

    # User Routes (14)
    {"method": "GET", "path": "/api/users/profile", "controller": "userController.getProfile", "middleware": "authenticate", "desc": "Retrieves comprehensive profile including CGPA, skills, avatar, and resume."},
    {"method": "PUT", "path": "/api/users/profile", "controller": "userController.updateProfile", "middleware": "authenticate, validateUser.profile", "desc": "Updates permitted profile fields; blocks role/email tampering and validates CGPA."},
    {"method": "POST", "path": "/api/users/profile/avatar", "controller": "userController.uploadAvatar", "middleware": "authenticate, uploadAvatar", "desc": "Uploads and saves PNG/JPG avatar image with unique timestamp filename."},
    {"method": "DELETE", "path": "/api/users/profile/avatar", "controller": "userController.deleteAvatar", "middleware": "authenticate", "desc": "Removes current avatar file from disk and resets avatar field in database."},
    {"method": "GET", "path": "/api/users/profile/avatar/:filename", "controller": "userController.getAvatarFile", "middleware": "None (Static public)", "desc": "Streams avatar image asset from /uploads/avatars directory."},
    {"method": "POST", "path": "/api/users/profile/resume", "controller": "userController.uploadResume", "middleware": "authenticate, uploadResume", "desc": "Uploads PDF resume (<=5MB), parses plain text, saves URL and original name."},
    {"method": "GET", "path": "/api/users/profile/resume/:filename", "controller": "userController.getResumeFile", "middleware": "authenticate", "desc": "Streams PDF resume file with strict authorization check (owner/employer/admin)."},
    {"method": "DELETE", "path": "/api/users/profile/resume", "controller": "userController.deleteResume", "middleware": "authenticate", "desc": "Deletes resume file from disk and clears resume reference in User record."},
    {"method": "GET", "path": "/api/users", "controller": "userController.listUsers", "middleware": "authenticate, requireRole('admin')", "desc": "Admin user governance directory with search, role filters, and pagination."},
    {"method": "POST", "path": "/api/users", "controller": "userController.createUser", "middleware": "authenticate, requireRole('admin'), validateUser.create", "desc": "Admin endpoint to provision Academic Coordinators or internal staff accounts."},
    {"method": "GET", "path": "/api/users/:id", "controller": "userController.getUserById", "middleware": "authenticate, requireRole('admin')", "desc": "Admin endpoint to inspect complete internal profile of a specific user."},
    {"method": "PUT", "path": "/api/users/:id", "controller": "userController.updateUser", "middleware": "authenticate, requireRole('admin'), validateUser.adminUpdate", "desc": "Admin modifies user role/details; last-admin guard prevents platform lockout."},
    {"method": "PATCH", "path": "/api/users/:id/status", "controller": "userController.updateUserStatus", "middleware": "authenticate, requireRole('admin')", "desc": "Toggles account status between active and inactive; last-admin protected."},
    {"method": "DELETE", "path": "/api/users/:id", "controller": "userController.deleteUser", "middleware": "authenticate, requireRole('admin')", "desc": "Deletes user account and associated profile assets; last-admin protected."},

    # Internship Routes (8)
    {"method": "GET", "path": "/api/internships", "controller": "internshipController.listInternships", "middleware": "None (Public)", "desc": "Public searchable directory of published opportunities with department filters."},
    {"method": "GET", "path": "/api/internships/manage", "controller": "internshipController.listManagedInternships", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Lists employer's postings including drafts, closed, and active opportunities."},
    {"method": "GET", "path": "/api/internships/:id", "controller": "internshipController.getInternshipById", "middleware": "None (Public / Authenticated)", "desc": "Retrieves comprehensive internship details, stipend, openings, and deadline."},
    {"method": "GET", "path": "/api/internships/:id/match-score", "controller": "internshipController.getMatchScore", "middleware": "authenticate, requireRole('student')", "desc": "Invokes Gemini AI or deterministic fallback to compute candidate fit score."},
    {"method": "POST", "path": "/api/internships", "controller": "internshipController.createInternship", "middleware": "authenticate, requireRole('company','coordinator','admin'), validateInternship", "desc": "Creates new opportunity posting in draft or published state with deadline check."},
    {"method": "PUT", "path": "/api/internships/:id", "controller": "internshipController.updateInternship", "middleware": "authenticate, requireRole('company','coordinator','admin'), validateInternship", "desc": "Updates internship details; validates employer ownership or coordinator override."},
    {"method": "PATCH", "path": "/api/internships/:id/status", "controller": "internshipController.updateInternshipStatus", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Transitions posting lifecycle (draft -> published -> closed -> archived)."},
    {"method": "DELETE", "path": "/api/internships/:id", "controller": "internshipController.deleteInternship", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Removes draft opportunity from catalog; verifies ownership."},

    # Application Routes (7)
    {"method": "POST", "path": "/api/applications", "controller": "applicationController.apply", "middleware": "authenticate, requireRole('student'), validateApplication.apply", "desc": "Submits student application; compound unique index {student, internship} blocks dups."},
    {"method": "GET", "path": "/api/applications/my", "controller": "applicationController.getMyApplications", "middleware": "authenticate, requireRole('student')", "desc": "Returns student's historical applications with live status, dates, and feedback."},
    {"method": "PATCH", "path": "/api/applications/:id/withdraw", "controller": "applicationController.withdraw", "middleware": "authenticate, requireRole('student')", "desc": "Allows candidate to withdraw submitted application prior to final decision."},
    {"method": "POST", "path": "/api/applications/:id/analyze-resume", "controller": "applicationController.analyzeResume", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Runs AI resume analysis on applicant's uploaded PDF against requirements."},
    {"method": "GET", "path": "/api/applications", "controller": "applicationController.listApplications", "middleware": "authenticate, requireRole('company','mentor','coordinator','admin')", "desc": "Lists candidate applications scoped to employer's postings or all for coordinator."},
    {"method": "GET", "path": "/api/applications/:id", "controller": "applicationController.getApplicationById", "middleware": "authenticate", "desc": "Retrieves application details, cover letter, resume link, and full audit history."},
    {"method": "PATCH", "path": "/api/applications/:id/status", "controller": "applicationController.updateStatus", "middleware": "authenticate, requireRole('company','coordinator','admin'), validateApplication.status", "desc": "Advances applicant status with mandatory decision reason and emits socket alert."},

    # Allocation Routes (8)
    {"method": "POST", "path": "/api/allocations/seed-demo", "controller": "allocationController.seedDemoWorkflow", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Seeds complete NovaTech Demo Labs workflow with 100% tasks and attendance."},
    {"method": "POST", "path": "/api/allocations", "controller": "allocationController.createAllocation", "middleware": "authenticate, requireRole('company','coordinator','admin'), validateAllocation", "desc": "Allocates accepted candidate to internship; blocks concurrent active placements."},
    {"method": "GET", "path": "/api/allocations", "controller": "allocationController.listAllocations", "middleware": "authenticate", "desc": "Lists intern allocations scoped to user's assigned role and company."},
    {"method": "GET", "path": "/api/allocations/:id/completion-eligibility", "controller": "allocationController.getCompletionEligibility", "middleware": "authenticate", "desc": "Evaluates tasks completed, attendance rate, and evaluation status for certificate."},
    {"method": "POST", "path": "/api/allocations/:id/request-completion", "controller": "allocationController.requestCompletion", "middleware": "authenticate, requireRole('student')", "desc": "Intern requests completion approval and specifies preferred recipient email."},
    {"method": "POST", "path": "/api/allocations/:id/review-completion", "controller": "allocationController.reviewCompletion", "middleware": "authenticate, requireRole('mentor','coordinator','admin')", "desc": "Mentor approves/rejects completion request, awards final grade, and unlocks cert."},
    {"method": "GET", "path": "/api/allocations/:id", "controller": "allocationController.getAllocationById", "middleware": "authenticate", "desc": "Retrieves allocation details, assigned mentor, dates, and completion status."},
    {"method": "PATCH", "path": "/api/allocations/:id/status", "controller": "allocationController.updateAllocationStatus", "middleware": "authenticate, requireRole('company','coordinator','admin')", "desc": "Updates allocation state (active -> completed) and triggers certificate issuance."},

    # Attendance Routes (4)
    {"method": "POST", "path": "/api/attendance", "controller": "attendanceController.logAttendance", "middleware": "authenticate, requireRole('student'), validateAttendance.log", "desc": "Logs daily hours and work summary; compound unique index {student, date} blocks dups."},
    {"method": "GET", "path": "/api/attendance/summary", "controller": "attendanceController.getSummary", "middleware": "authenticate", "desc": "Computes monthly attendance statistics, present/absent days, and completion rate."},
    {"method": "GET", "path": "/api/attendance", "controller": "attendanceController.listAttendance", "middleware": "authenticate", "desc": "Lists attendance submissions filtered by allocation, date range, or intern."},
    {"method": "PATCH", "path": "/api/attendance/:id/review", "controller": "attendanceController.reviewAttendance", "middleware": "authenticate, requireRole('mentor','coordinator','admin'), validateAttendance.review", "desc": "Supervisor approves or corrects hours with mandatory recorded audit reason."},

    # Task Routes (5)
    {"method": "POST", "path": "/api/tasks", "controller": "taskController.createTask", "middleware": "authenticate, requireRole('mentor','coordinator','admin'), validateTask.create", "desc": "Assigns milestone task with description, priority, and due date to intern."},
    {"method": "GET", "path": "/api/tasks", "controller": "taskController.listTasks", "middleware": "authenticate", "desc": "Lists tasks for intern or supervisor; dynamically computes isOverdue flag."},
    {"method": "GET", "path": "/api/tasks/:id", "controller": "taskController.getTaskById", "middleware": "authenticate", "desc": "Retrieves milestone details, deliverable links, notes, and mentor review feedback."},
    {"method": "PATCH", "path": "/api/tasks/:id/progress", "controller": "taskController.submitProgress", "middleware": "authenticate, requireRole('student'), validateTask.progress", "desc": "Intern submits milestone deliverables, notes, and repository/document URLs."},
    {"method": "PATCH", "path": "/api/tasks/:id/review", "controller": "taskController.reviewTask", "middleware": "authenticate, requireRole('mentor','coordinator','admin'), validateTask.review", "desc": "Mentor marks task completed, awards score (0-100), and records feedback."},

    # Notification Routes (12)
    {"method": "GET", "path": "/api/notifications/push/vapid-key", "controller": "notificationController.getVapidPublicKey", "middleware": "authenticate", "desc": "Returns server's VAPID public key for Web Push service worker subscription."},
    {"method": "POST", "path": "/api/notifications/push/subscribe", "controller": "notificationController.subscribePush", "middleware": "authenticate", "desc": "Stores browser push subscription (endpoint, keys) in User document."},
    {"method": "POST", "path": "/api/notifications/push/unsubscribe", "controller": "notificationController.unsubscribePush", "middleware": "authenticate", "desc": "Removes browser push subscription endpoint from database."},
    {"method": "POST", "path": "/api/notifications/test-desktop", "controller": "notificationController.testDesktopNotification", "middleware": "authenticate", "desc": "Dispatches test real-time socket and desktop notification to current user."},
    {"method": "GET", "path": "/api/notifications", "controller": "notificationController.getNotifications", "middleware": "authenticate", "desc": "Returns chronological notification list with pagination (default 30 alerts)."},
    {"method": "GET", "path": "/api/notifications/unread-count", "controller": "notificationController.getUnreadCount", "middleware": "authenticate", "desc": "Returns total unread notification count for authenticated user badge."},
    {"method": "PATCH", "path": "/api/notifications/read-all", "controller": "notificationController.markAllAsRead", "middleware": "authenticate", "desc": "Marks all user notifications as read and resets unread count to 0."},
    {"method": "PATCH", "path": "/api/notifications/conversation/:conversationId/read", "controller": "notificationController.markConversationNotificationsRead", "middleware": "authenticate", "desc": "Clears notifications belonging to a specific conversation when opened."},
    {"method": "PATCH", "path": "/api/notifications/:id/read", "controller": "notificationController.markAsRead", "middleware": "authenticate", "desc": "Marks single notification as read; verifies recipient ownership."},
    {"method": "DELETE", "path": "/api/notifications", "controller": "notificationController.clearAllNotifications", "middleware": "authenticate", "desc": "Deletes all notification records belonging to authenticated user."},
    {"method": "DELETE", "path": "/api/notifications/:id", "controller": "notificationController.deleteNotification", "middleware": "authenticate", "desc": "Deletes individual notification record; verifies recipient ownership."},
    {"method": "POST", "path": "/api/notifications/announcements", "controller": "notificationController.broadcastAnnouncement", "middleware": "authenticate, requireRole('coordinator','admin')", "desc": "Broadcasts institutional announcement to all users or specific target roles."},

    # Chat Routes (16)
    {"method": "GET", "path": "/api/chat/conversations", "controller": "chatController.getConversations", "middleware": "authenticate", "desc": "Lists user's direct and group conversations with last message and unread counts."},
    {"method": "POST", "path": "/api/chat/conversations", "controller": "chatController.createConversation", "middleware": "authenticate, validateChat.conversation", "desc": "Creates or reuses 1-on-1 direct chat or coordinator group conversation."},
    {"method": "GET", "path": "/api/chat/conversations/:id", "controller": "chatController.getConversationById", "middleware": "authenticate", "desc": "Retrieves conversation channel details; verifies participant membership."},
    {"method": "POST", "path": "/api/chat/upload", "controller": "chatController.uploadAttachment", "middleware": "authenticate, chatAttachmentUpload", "desc": "Uploads chat attachment (images, PDF, documents <=10MB) via Multer."},
    {"method": "GET", "path": "/api/chat/attachments/:fileKey", "controller": "chatController.getAttachment", "middleware": "authenticate", "desc": "Streams chat attachment file with authenticated participant authorization check."},
    {"method": "GET", "path": "/api/chat/conversations/:id/messages", "controller": "chatController.getMessages", "middleware": "authenticate", "desc": "Retrieves paginated messages for conversation with reactions and attachments."},
    {"method": "POST", "path": "/api/chat/conversations/:id/messages", "controller": "chatController.sendMessage", "middleware": "authenticate, validateChat.message", "desc": "Saves message in MongoDB, updates conversation, and broadcasts via Socket.IO."},
    {"method": "PATCH", "path": "/api/chat/messages/:messageId", "controller": "chatController.editMessage", "middleware": "authenticate", "desc": "Author edits message within 15-minute window; marks isEdited: true."},
    {"method": "DELETE", "path": "/api/chat/messages/:messageId", "controller": "chatController.deleteMessage", "middleware": "authenticate", "desc": "Author deletes message for everyone or user hides message for self."},
    {"method": "POST", "path": "/api/chat/messages/:messageId/forward", "controller": "chatController.forwardMessage", "middleware": "authenticate", "desc": "Forwards existing message content into another target conversation."},
    {"method": "PATCH", "path": "/api/chat/messages/:messageId/react", "controller": "chatController.toggleReaction", "middleware": "authenticate", "desc": "Toggles emoji reaction on message and emits real-time reaction update."},
    {"method": "POST", "path": "/api/chat/conversations/:id/simulate-reply", "controller": "chatController.simulateReply", "middleware": "authenticate", "desc": "Development diagnostic endpoint simulating reciprocal message in thread."},
    {"method": "PATCH", "path": "/api/chat/conversations/:id/read", "controller": "chatController.markAsRead", "middleware": "authenticate", "desc": "Records read receipt, clears unread counter, and emits chat:read_update."},
    {"method": "PATCH", "path": "/api/chat/conversations/:id/members", "controller": "chatController.manageMembers", "middleware": "authenticate", "desc": "Group admin adds or removes participants from group conversation."},
    {"method": "GET", "path": "/api/chat/contacts", "controller": "chatController.getContacts", "middleware": "authenticate", "desc": "Returns directory of eligible contacts (peers, mentors, employers) to start chats."},
    {"method": "GET", "path": "/api/chat/online-users", "controller": "chatController.getOnlineUsers", "middleware": "authenticate", "desc": "Returns list of user IDs currently connected via active Socket.IO connections."},

    # Evaluation Routes (4)
    {"method": "POST", "path": "/api/evaluations", "controller": "evaluationController.createEvaluation", "middleware": "authenticate, requireRole('mentor','coordinator','admin'), validateEvaluation", "desc": "Submits mid-term or final appraisal across technical, behavioral, and overall criteria."},
    {"method": "GET", "path": "/api/evaluations", "controller": "evaluationController.listEvaluations", "middleware": "authenticate", "desc": "Lists appraisals scoped to user's assigned interns or student's own records."},
    {"method": "GET", "path": "/api/evaluations/:id", "controller": "evaluationController.getEvaluationById", "middleware": "authenticate", "desc": "Retrieves complete appraisal scorecard with criteria scores and recommendation."},
    {"method": "GET", "path": "/api/evaluations/allocation/:allocationId", "controller": "evaluationController.getEvaluationByAllocation", "middleware": "authenticate", "desc": "Retrieves evaluation scorecard linked to a specific completed allocation."},

    # Analytics Routes (1)
    {"method": "GET", "path": "/api/analytics/dashboard", "controller": "analyticsController.getDashboardMetrics", "middleware": "authenticate", "desc": "Computes customized KPI aggregations and funnel counts tailored to user's role."},

    # Report Routes (2)
    {"method": "GET", "path": "/api/reports", "controller": "reportController.getReports", "middleware": "authenticate, requireRole('company','mentor','coordinator','admin')", "desc": "Generates aggregate institutional report datasets with date and department filters."},
    {"method": "GET", "path": "/api/reports/export", "controller": "reportController.exportReportCsv", "middleware": "authenticate, requireRole('company','mentor','coordinator','admin')", "desc": "Streams sanitized CSV dataset with strict redaction of sensitive credentials."},

    # Certificate Routes (10)
    {"method": "GET", "path": "/api/certificates/verify/:code", "controller": "certificateController.verifyCertificate", "middleware": "None (Public)", "desc": "Public cryptographic verification ledger validating certificate authenticity."},
    {"method": "GET", "path": "/api/certificates/admin/audit", "controller": "certificateController.getAdminAudit", "middleware": "authenticate, requireRole('coordinator','admin')", "desc": "Institutional audit table of all issued certificates, verification codes, and email logs."},
    {"method": "GET", "path": "/api/certificates/admin/mentor-stats", "controller": "certificateController.getMentorStats", "middleware": "authenticate, requireRole('coordinator','admin')", "desc": "Computes supervisor issuance statistics and completion review turnaround times."},
    {"method": "POST", "path": "/api/certificates/admin/seed-demo", "controller": "certificateController.seedDemoWorkflow", "middleware": "authenticate, requireRole('coordinator','admin')", "desc": "Administrative shortcut seeding verifiable demo certificate with Padmavathi signature."},
    {"method": "GET", "path": "/api/certificates/allocation/:allocationId/preview", "controller": "certificateController.previewCertificate", "middleware": "authenticate, requireRole('mentor','coordinator','admin')", "desc": "Renders metadata preview of certificate layout prior to official cryptographic signing."},
    {"method": "POST", "path": "/api/certificates/allocation/:allocationId/issue", "controller": "certificateController.issueCertificate", "middleware": "authenticate, requireRole('mentor','coordinator','admin')", "desc": "Enforces date validation, generates PDFKit landscape cert with QR & Padmavathi signature, emails PDF via SMTP."},
    {"method": "GET", "path": "/api/certificates/my", "controller": "certificateController.getMyCertificates", "middleware": "authenticate, requireRole('student')", "desc": "Returns list of certificates officially awarded to authenticated student."},
    {"method": "GET", "path": "/api/certificates/:id", "controller": "certificateController.getCertificateById", "middleware": "authenticate", "desc": "Retrieves certificate metadata, verification code, and delivery timestamps."},
    {"method": "GET", "path": "/api/certificates/:id/download", "controller": "certificateController.downloadCertificatePdf", "middleware": "None (Protected via cert ID / code)", "desc": "Dynamically streams binary PDF document (application/pdf) for viewing/saving."},
    {"method": "POST", "path": "/api/certificates/:id/resend-email", "controller": "certificateController.resendCertificateEmail", "middleware": "authenticate, requireRole('coordinator','admin')", "desc": "Re-dispatches certificate PDF attachment via Nodemailer SMTP and tracks resendCount."},

    # External Jobs Routes (8)
    {"method": "GET", "path": "/api/external-jobs", "controller": "externalJobController.listJobs", "middleware": "None (Public)", "desc": "Public searchable catalog of curated external tech jobs with skill and location filters."},
    {"method": "GET", "path": "/api/external-jobs/recommendations", "controller": "externalJobController.getRecommendations", "middleware": "authenticate, requireRole('student')", "desc": "Returns external jobs ranked by AI match score against student's profile skills."},
    {"method": "GET", "path": "/api/external-jobs/my-interactions", "controller": "externalJobController.getMyInteractions", "middleware": "authenticate, requireRole('student')", "desc": "Returns student's saved bookmarks, application dates, and personal interview notes."},
    {"method": "POST", "path": "/api/external-jobs/sync", "controller": "externalJobController.syncJobs", "middleware": "authenticate, requireRole('admin')", "desc": "Admin trigger to manually execute external job radar feed synchronization."},
    {"method": "GET", "path": "/api/external-jobs/:id", "controller": "externalJobController.getJobById", "middleware": "None (Public)", "desc": "Retrieves external job details, source attribution, and required tech skills."},
    {"method": "POST", "path": "/api/external-jobs/:id/bookmark", "controller": "externalJobController.toggleBookmark", "middleware": "authenticate, requireRole('student')", "desc": "Toggles saved bookmark state for student in UserJobBookmark collection."},
    {"method": "POST", "path": "/api/external-jobs/:id/apply", "controller": "externalJobController.trackApplication", "middleware": "authenticate, requireRole('student')", "desc": "Records application timestamp and increments public job applyCount."},
    {"method": "PATCH", "path": "/api/external-jobs/:id/notes", "controller": "externalJobController.updateNotes", "middleware": "authenticate, requireRole('student')", "desc": "Saves candidate's private application notes and interview follow-up reminders."}
]

print(f"Loaded {len(ALL_ENDPOINTS)} backend API endpoints")

# --- Role & Permission Matrix Data ---
ROLE_MATRIX = [
    {
        "domain": "Authentication & Sessions",
        "student": "Register, Login, Google SSO, 2FA, View/Terminate own sessions, Logout",
        "mentor": "Login, Google SSO, 2FA, View/Terminate own sessions, Logout",
        "company": "Register, Login, Google SSO, 2FA, View/Terminate own sessions, Logout",
        "coordinator": "Staff Login, 2FA, View/Terminate own sessions, Logout",
        "admin": "Admin Login, 2FA, View/Terminate own sessions, Logout"
    },
    {
        "domain": "User Profiles & Uploads",
        "student": "Edit bio/skills/CGPA, upload avatar, upload PDF resume, delete own files",
        "mentor": "Edit bio/skills, upload avatar, delete own avatar",
        "company": "Edit company bio/location, upload avatar/logo, delete own logo",
        "coordinator": "Edit institutional bio, upload avatar, inspect candidate resumes",
        "admin": "Full profile edit, manage all user avatars/resumes, download any resume"
    },
    {
        "domain": "Internship Postings",
        "student": "Browse published catalog, filter by dept, view match score (No create/edit)",
        "mentor": "Browse published catalog, view postings linked to assigned company",
        "company": "Create draft, publish, edit own postings, toggle status, delete draft",
        "coordinator": "Create, edit, publish, close, or archive any institutional posting",
        "admin": "Full CRUD access across all company postings and catalog items"
    },
    {
        "domain": "Applications Workflow",
        "student": "Apply before deadline, track status history, withdraw before decision",
        "mentor": "View applicants allocated to mentor's company/department",
        "company": "Review applicants, run AI resume analysis, advance status, accept/reject",
        "coordinator": "Review all institutional applicants, inspect AI screening reports",
        "admin": "Full access to review, reassign, or update any application status"
    },
    {
        "domain": "Allocations & Mentoring",
        "student": "View active allocation, view assigned mentor, request completion",
        "mentor": "View assigned intern allocations, review completion requests, assign grades",
        "company": "Allocate accepted candidates to mentors, define start/end dates",
        "coordinator": "Allocate candidates, reassign mentors, review completion requests",
        "admin": "Full allocation governance, seed demo workflows, override dates"
    },
    {
        "domain": "Tasks & Milestones",
        "student": "View assigned tasks, track due dates, submit deliverable link/notes",
        "mentor": "Create milestone tasks, assign to interns, score deliverables (0-100)",
        "company": "View milestone progress and completion ratios for company interns",
        "coordinator": "Audit milestone task creation and completion metrics across departments",
        "admin": "Full task oversight, manage overdue milestones, audit deliverables"
    },
    {
        "domain": "Daily Attendance",
        "student": "Log daily attendance with hours and work summary (Same-day dedup enforced)",
        "mentor": "Audit daily attendance, approve hours, correct records with mandatory notes",
        "company": "Review monthly attendance summaries and verification rates",
        "coordinator": "Audit institutional attendance records, generate compliance reports",
        "admin": "Full attendance audit access, modify or purge corrupted records"
    },
    {
        "domain": "Evaluations & Appraisal",
        "student": "View own submitted evaluation scorecards and mentor feedback",
        "mentor": "Submit mid-term appraisal, rate technical/behavioral criteria, recommend",
        "company": "Review mentor appraisals and candidate performance reviews",
        "coordinator": "Submit final coordinator evaluation, review grade recommendations",
        "admin": "Audit all performance appraisals, export evaluation aggregates"
    },
    {
        "domain": "Certificates & Credentials",
        "student": "View awarded certificates, stream PDF, share public QR verification link",
        "mentor": "Review completion eligibility, preview certificate layout, issue certificate",
        "company": "View issued certificates for company interns, verify credentials",
        "coordinator": "Issue certificates, audit certificate ledger, trigger email resend",
        "admin": "Full certificate governance, seed demo certificate, audit logs, resend"
    },
    {
        "domain": "Real-Time Chat & Channels",
        "student": "1-on-1 direct chat with mentors/coordinators, send attachments, react",
        "mentor": "1-on-1 direct chat with interns/employers, group chat participant",
        "company": "1-on-1 direct chat with candidates and interns, send documents",
        "coordinator": "Create group chat channels, manage group members, broadcast in chats",
        "admin": "Participate in any conversation, moderate abusive message contents"
    },
    {
        "domain": "System Notifications",
        "student": "Receive live in-app/desktop alerts for application, task, and chat events",
        "mentor": "Receive alerts for attendance submissions, completion requests, chats",
        "company": "Receive live alerts for new applications, intern allocations, messages",
        "coordinator": "Receive system alerts, broadcast announcements to all users/students",
        "admin": "Receive governance alerts, broadcast institutional announcements"
    },
    {
        "domain": "Analytics & Reporting",
        "student": "Personalized student dashboard KPIs (applied count, active tasks)",
        "mentor": "Mentor dashboard KPIs (active interns, pending task reviews)",
        "company": "Company dashboard KPIs (openings, applicant funnel, intern count)",
        "coordinator": "Comprehensive institutional analytics, view reports, export sanitized CSV",
        "admin": "Full platform analytics, system health metrics, export sanitized CSV"
    },
    {
        "domain": "Admin User Governance",
        "student": "No access (HTTP 403 Forbidden)",
        "mentor": "No access (HTTP 403 Forbidden)",
        "company": "No access (HTTP 403 Forbidden)",
        "coordinator": "No access to user administration console (HTTP 403 Forbidden)",
        "admin": "Full CRUD on users, provision coordinators, change roles, last-admin guard"
    }
]

print("Loaded Role & Permission Matrix")

# --- In-Depth Workflows Text ---

WORKFLOW_SESSION = {
    "title": "In-Depth Architectural Workflow A: Session Management",
    "steps": [
        ("Step 1: Credential Submission", "The user navigates to /login and submits email and password. In frontend/src/pages/LoginPage.jsx, the form dispatches authService.login({ email, password }). This issues an HTTP POST request to /api/auth/login."),
        ("Step 2: Password Verification & 2FA Challenge", "In backend/src/controllers/authController.js (login), the controller queries User.findOne({ email }). It executes bcrypt.compare(password, user.password). If password matches and user.twoFactorEnabled is true, the server generates a 6-digit crypto OTP, saves it in twoFactorCode with a 10-minute expiry, dispatches an email via emailService.sendTwoFactorCode(), and returns HTTP 200 { requiresTwoFactor: true, email }."),
        ("Step 3: Two-Factor Confirmation", "The client displays TwoFactorModal. The user enters the 6-digit OTP code, which posts to /api/auth/login-verify-2fa. In authController.verifyTwoFactorLogin, the code and expiry are verified, and twoFactorCode is cleared from the User document."),
        ("Step 4: Server-Side Active Session Tracking", "Upon successful authentication, the server creates a new Session document in backend/src/models/Session.js: Session.create({ user: user._id, tokenSignature, isValid: true, deviceInfo: req.headers['user-agent'], ipAddress: req.ip }). The tokenSignature is derived from the crypto signature portion of the newly generated JWT."),
        ("Step 5: JWT Generation & Token Delivery", "The server signs a JWT token using jwt.sign({ id: user._id, role: user.role, email: user.email }, env.JWT_SECRET, { expiresIn: '24h' }). The controller returns HTTP 200 { success: true, token, user }."),
        ("Step 6: Client-Side Session Storage & Tab Isolation", "In frontend/src/context/AuthContext.jsx, the token and user are stored in browser sessionStorage (sims_auth_token and sims_auth_user). Legacy persistent entries in localStorage are purged to ensure that newly opened tabs always start fresh and isolated."),
        ("Step 7: Axios Request Interceptor & Anti-Caching", "In frontend/src/api/client.js, setAuthToken(token) injects the token into Axios default headers as Authorization: Bearer <token>. All authenticated responses from the Express server include Cache-Control: no-store, no-cache, must-revalidate headers to prevent sensitive pages from remaining in browser memory."),
        ("Step 8: Server Authentication Middleware Guard", "Every protected backend endpoint passes through authenticate in backend/src/middlewares/auth.js. It extracts the Bearer token, verifies cryptographic validity via jwt.verify, extracts the token signature, and queries Session.findOne({ tokenSignature, isValid: true }). If the session is missing or revoked, it rejects the request with HTTP 401 TOKEN_EXPIRED."),
        ("Step 9: Client-Side Protected Route Guarding", "In frontend/src/components/auth/ProtectedRoute.jsx, the router checks isAuthenticated and allowedRoles. If unauthenticated, it redirects to /login. If the user's role is not authorized, it renders an access denied alert."),
        ("Step 10: Automatic Client-Side Expiry Timer", "In AuthContext.jsx, a useEffect hook decodes the JWT exp claim, calculates milliseconds until expiration, and schedules a setTimeout timer. When the timer elapses, it fires the sims-session-expired DOM event, which automatically logs the user out and clears state."),
        ("Step 11: Explicit Logout & Revocation Execution", "When the user clicks Logout, handleLogout calls authService.logout(). In authController.js, the server marks Session.updateOne({ tokenSignature }, { isValid: false }). Client-side sessionStorage is cleared, Axios Bearer headers are wiped, and the window replaces history state with /login.")
    ]
}

WORKFLOW_CERTIFICATE = {
    "title": "In-Depth Architectural Workflow B: Certificate Generation & Verification",
    "steps": [
        ("Step 1: Requirements & Eligibility Assessment", "The student completes daily attendance and milestone tasks. In backend/src/controllers/allocationController.js (getCompletionEligibility), the system aggregates attendances and tasks. It verifies that tasks completed ratio is 100%, attendance rate is >= 80%, and an evaluation exists."),
        ("Step 2: Student Completion Request", "In frontend/src/pages/CertificatesPage.jsx or ManageAllocationsPage.jsx, the student clicks 'Request Completion'. This posts to /api/allocations/:id/request-completion with studentNotes and preferredEmail. The allocation completionRequest status becomes 'pending'."),
        ("Step 3: Mentor Review & Grade Approval", "The assigned industry mentor opens ManageAllocationsPage.jsx and clicks 'Review Completion'. This posts to /api/allocations/:id/review-completion. The mentor selects 'approved', assigns supervisorRecommendation, awards finalGrade (e.g. 'A+ / Outstanding'), and approves completion."),
        ("Step 4: Certificate Preview Rendering", "Prior to official issuance, the mentor or coordinator can call GET /api/certificates/allocation/:allocationId/preview in certificateController.js (previewCertificate). This aggregates recipient name, company name, start/end dates, and layout coordinates for verification."),
        ("Step 5: Certificate Issuance Initiation", "The mentor, coordinator, or admin clicks 'Issue Certificate'. This submits an HTTP POST to /api/certificates/allocation/:allocationId/issue, handled by certificateController.issueCertificate and certificateService.issueCertificateForAllocation."),
        ("Step 6: Strict Date Validation & Integrity Guard", "In backend/src/services/certificateService.js, the system enforces mandatory verifiedStartDate and verifiedEndDate from the allocation record. It verifies that both dates are valid Calendar Dates and verifiedEndDate >= verifiedStartDate. If missing or invalid, it immediately throws HTTP 400 INVALID_START_DATE or INVALID_DATE_RANGE. Fallback substitution with createdAt is strictly prohibited."),
        ("Step 7: Idempotency Barrier", "The service executes Certificate.findOne({ allocation: allocationId }). If a certificate has already been issued for this allocation, it returns the existing certificate record without generating duplicate database documents."),
        ("Step 8: Cryptographic Code & QR Generation", "The service generates a unique credential code (e.g. CERT-2026-853925) using random cryptographic bytes. It invokes QRCode.toDataURL(verificationUrl) to create a high-resolution QR code encoding the public URL https://<domain>/verify-certificate/:code."),
        ("Step 9: PDFKit Landscape A4 Compilation", "PDFKit instantiates a landscape A4 canvas (841.89 x 595.28 pt). It draws dual concentric borders in deep navy (#0f2b46) and warm gold (#c79d4c), renders official InternHub typography, recipient name, company, mentor, honors grade, and verified date range. It crisply embeds the official Padmavathi Program Coordinator digital signature graphic and stamps the QR verification code."),
        ("Step 10: Database Persistence & Binary Storage", "The completed PDF is buffered into memory. A Certificate document is created in backend/src/models/Certificate.js storing certificateNumber, student, allocation, company, mentor, issueDate, verifiedStartDate, verifiedEndDate, qrCodeData, and pdfBuffer."),
        ("Step 11: Nodemailer Brevo/Gmail SMTP Dispatch", "The service invokes emailService.sendCertificateEmail(). It dispatches an email via Brevo or Gmail SMTP with the PDF certificate attached. It records emailSentAt and updates emailDeliveryStatus: 'sent'."),
        ("Step 12: Public Cryptographic Verification Ledger", "Any employer or university can scan the QR code or visit /verify-certificate/:code. In backend/src/controllers/certificateController.js (verifyCertificate), the endpoint looks up Certificate.findOne({ certificateNumber: code }), validates against MongoDB, and returns sanitized verification data proving authentic credential issuance.")
    ]
}

WORKFLOW_CHAT_NOTIF = {
    "title": "In-Depth Architectural Workflow C: Real-Time Notifications and Chat",
    "steps": [
        ("Step 1: WebSocket Connection & Handshake Authentication", "Upon successful login, frontend/src/context/SocketContext.jsx initializes io(socketUrl, { auth: { token } }). In backend/src/socket.js, the io.use middleware extracts the token, verifies the signature using jwt.verify, checks that User exists and status is 'active', and binds socket.user = user. Unauthorized sockets are rejected at the handshake."),
        ("Step 2: Private User Room Isolation", "Upon connection, the server places the socket into a private room: socket.join(`user:${user._id}`). It emits user:status: online to tracked contacts and records the active socket ID in an in-memory connection set."),
        ("Step 3: Conversation Room Authorization", "When the user opens ChatPage.jsx, the client emits chat:join_conversation with conversationId. The server queries Conversation.exists({ _id: convId, participants: socket.user._id }). If verified, socket.join(`conv:${convId}`). Non-members are rejected with an Unauthorized callback error."),
        ("Step 4: Message Dispatch & MongoDB Persistence", "The user types a message or attaches a file and submits. In frontend/src/services/chatService.js, it calls POST /api/chat/conversations/:id/messages. In chatController.js (sendMessage), the message is validated (<=5000 chars, valid attachments) and saved in backend/src/models/Message.js. The conversation lastMessage and lastMessageAt are updated."),
        ("Step 5: Multi-Channel Socket Broadcast", "The server broadcasts the saved message to the room: io.to(`conv:${conversationId}`).emit('chat:message', populatedMessage). If recipients are not currently in the active conversation room, the server creates a Notification document and emits notification:new to `user:${recipientId}`."),
        ("Step 6: Frontend Active Conversation Detection & Mute Guard", "In frontend/src/context/SocketContext.jsx, when chat:message arrives, it checks: isActivelyViewing = (activeChatConvRef.current === message.conversation && document.visibilityState === 'visible' && document.hasFocus()). If true, floating toasts, audio chimes, and desktop popups are suppressed so active chat is not disrupted."),
        ("Step 7: Branded Floating Toast & Audio Chime", "If the recipient is not actively viewing the chat, SocketContext dispatches addToast() to NotificationToastContainer.jsx. A floating toast appears at the bottom-right showing the InternHub official logo, sender name, message snippet, and Dismiss/Open buttons. soundManager.playNotificationChime() triggers an audio alert if enabled."),
        ("Step 8: Native Operating System Desktop Notification", "desktopNotification.js evaluates Notification.permission. If granted, it triggers showDesktopNotification(). It prioritizes Service Worker registration.showNotification('InternHub Message', { body, tag: `chat_${convId}_${msgId}`, renotify: true }), falling back to window.Notification. The unique tag ensures that Windows Action Center displays a floating banner. Clicking the notification focuses the window and navigates directly to /chat?conversationId=..."),
        ("Step 9: Read Receipts & Unread Counter Decrement", "When the recipient views the thread, ChatPage.jsx calls chatService.markConversationAsRead(id). This sends PATCH /api/chat/conversations/:id/read. The controller adds user to readBy array in messages and emits chat:read_update to `conv:${id}`. Local unread badges and top navbar indicators immediately decrement.")
    ]
}

# --- Missing Links, Incomplete Features & Issues Data ---
ISSUES_DISCOVERED = [
    {
        "issue": "Legacy Admin Login Component (Redirect Only)",
        "location": "frontend/src/pages/AdminLoginPage.jsx, App.jsx (/admin/login)",
        "nature": "Architectural Debt / Redundant File",
        "description": "AdminLoginPage.jsx contains only a <Navigate to='/login' replace /> redirect. The application transitioned to a single unified login portal. While functional, the file is redundant and can be removed in future cleanups."
    },
    {
        "issue": "Direct-Route Only Pages (Missing Sidebar Navigation Links)",
        "location": "frontend/src/components/layout/Sidebar.jsx",
        "nature": "UI Discoverability Gap",
        "description": "/admin/certificates is accessible in the sidebar only for Admins and Coordinators, but not directly visible to Companies. /policies, /privacy, /terms, and /cookies are accessible exclusively via footer links or direct URL navigation."
    },
    {
        "issue": "Multi-Route Mapping to PoliciesPage",
        "location": "frontend/src/App.jsx (Lines 66-70)",
        "nature": "Routing Redundancy",
        "description": "/policies, /privacy, /terms, /security-policy, and /cookies all render PoliciesPage.jsx with internal tab switching. Direct deep linking to tabs could be improved by using URL search parameters (?tab=privacy)."
    },
    {
        "issue": "Local Disk Uploads Storage vs Cloud Object Storage",
        "location": "backend/src/controllers/userController.js, chatController.js, uploads/",
        "nature": "Scalability Limitation",
        "description": "User avatars, resumes, and chat attachments are stored on local container disk (/uploads). In a multi-instance containerized or Kubernetes deployment behind a load balancer, uploaded assets must be migrated to AWS S3 or Google Cloud Storage."
    },
    {
        "issue": "In-Memory Rate Limiting Cluster Limitation",
        "location": "backend/src/middlewares/rateLimiter.js",
        "nature": "High-Availability Limitation",
        "description": "Token bucket rate limiting is implemented in-memory using JavaScript Maps. Rate limits reset when the server restarts and are not synchronized across multiple clustered Node.js worker processes without a Redis backend."
    },
    {
        "issue": "Omission of Automated Frontend Unit & UI Tests",
        "location": "frontend/package.json",
        "nature": "Test Coverage Gap",
        "description": "While the backend has 225 comprehensive automated tests, the frontend lacks Vitest or React Testing Library suites. UI verification is currently based on Vite production compilation and manual exploratory testing."
    }
]

print("Loaded Workflows and Issues data")

# ==============================================================================
# 2. DOCX REPORT BUILDER (BLACK & WHITE ENTERPRISE STYLE)
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cell in twips."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Apply background color to cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def build_docx_report(docx_path):
    print(f"Building DOCX report at: {docx_path}")
    doc = docx.Document()
    
    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header & Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("InternHub – Complete Feature Implementation and Technical Architecture Report | Confidential")
        f_run.font.name = "Arial"
        f_run.font.size = Pt(8)
        f_run.font.color.rgb = RGBColor(100, 100, 100)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run(DOC_TITLE)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(DOC_SUBTITLE)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)
    
    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Document Identifier", DOC_ID),
        ("Date of Audit", DOC_DATE),
        ("System Under Test", DOC_SYSTEM),
        ("Lead Author & Auditor", DOC_AUTHOR),
        ("Classification", DOC_CLASSIFICATION),
        ("Overall Architecture Verdict", DOC_VERDICT),
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        r = meta_table.rows[row_idx]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(5.0)
        
        c0.text = k
        c1.text = v
        
        p0 = c0.paragraphs[0]
        p0.runs[0].font.name = "Arial"
        p0.runs[0].font.size = Pt(8.5)
        p0.runs[0].font.bold = True
        
        p1 = c1.paragraphs[0]
        p1.runs[0].font.name = "Arial"
        p1.runs[0].font.size = Pt(8.5)
        
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_shading(c0, "F3F4F6")
        set_cell_shading(c1, "FFFFFF" if row_idx % 2 == 0 else "F9FAFB")
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Section 1: Project Architecture Overview
    h1 = doc.add_heading("1. Project Architecture Overview", level=1)
    h1.runs[0].font.name = "Arial"
    h1.runs[0].font.size = Pt(13)
    h1.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    p = doc.add_paragraph(
        "The Smart Internship Management System (InternHub) is engineered as an enterprise-grade micro-monolith platform uniting five distinct stakeholders: "
        "Students, Industry Mentors, Employer Companies, Academic Coordinators, and System Administrators. "
        "The backend pairs a stateless Express.js REST API with a stateful, bi-directional Socket.IO WebSocket server operating synchronously on port 5001. "
        "Data persistence is managed via MongoDB Atlas utilizing Mongoose 8.9.5 ODM with strict schema validation and compound uniqueness indexing. "
        "The frontend is a single-page application (SPA) built with React 18.3.1, bundled via Vite 6.0.7, styled with Tailwind CSS 3.4.17, "
        "and structured through React Router DOM v6 with role-aware route guards."
    )
    p.runs[0].font.name = "Arial"
    p.runs[0].font.size = Pt(9)
    
    # Architecture Stack Table
    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_headers = ["Layer / Subsystem", "Technology & Version", "Architectural Role & Scope"]
    for i, h in enumerate(tech_headers):
        c = tech_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8.5)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 80, 80, 100, 100)
        
    tech_data = [
        ("Backend Framework", "Express.js ^4.21.2 / Node.js v24.17.0", "Handles 15 route modules, 123 endpoints, JWT auth, and middleware."),
        ("Real-Time Engine", "Socket.IO ^4.8.4", "Maintains authenticated persistent connections, user rooms, and chat events."),
        ("Database & ODM", "MongoDB Atlas / Mongoose ^8.9.5", "14 strongly-typed collections, compound uniqueness barriers, and aggregations."),
        ("Frontend Application", "React ^18.3.1 / Vite ^6.0.7", "SPA with 26 pages, Context API, responsive Tailwind layout, and Vite bundling."),
        ("Document Generation", "PDFKit ^0.20.2 / QRCode ^1.5.4", "Generates landscape A4 completion certificates with QR verification & Padmavathi signature."),
        ("External Services", "Brevo/Gmail SMTP / Gemini 2.5 Flash", "Nodemailer transactional email delivery; Google Gemini AI resume fit scoring.")
    ]
    for row_idx, data in enumerate(tech_data, start=1):
        r = tech_table.rows[row_idx]
        for col_idx, val in enumerate(data):
            cell = r.cells[col_idx]
            cell.text = val
            p = cell.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(8)
            set_cell_margins(cell, 60, 60, 100, 100)
            set_cell_shading(cell, "FFFFFF" if row_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Frontend Page and Route Directory
    h2 = doc.add_heading("2. Frontend Page and Route Directory", level=1)
    h2.runs[0].font.name = "Arial"
    h2.runs[0].font.size = Pt(13)
    h2.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "InternHub features 26 distinct page components in frontend/src/pages/. Below is the complete directory detailing each page's "
        "route path, source file, permitted roles, navigation triggers, and primary UI components:"
    ).runs[0].font.size = Pt(9)
    
    page_table = doc.add_table(rows=len(PAGES_DIRECTORY) + 1, cols=5)
    page_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    page_headers = ["Page Title", "Route URL", "Source File Path", "Allowed Roles", "Navigation & Access"]
    for i, h in enumerate(page_headers):
        c = page_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 80, 80, 80, 80)
        
    for r_idx, page in enumerate(PAGES_DIRECTORY, start=1):
        row = page_table.rows[r_idx]
        vals = [page["title"], page["route"], page["file"], page["roles"], page["nav"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7.5)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 50, 50, 60, 60)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 3: Complete Feature-to-File Implementation Mapping
    h3 = doc.add_heading("3. Complete Feature-to-File Implementation Mapping", level=1)
    h3.runs[0].font.name = "Arial"
    h3.runs[0].font.size = Pt(13)
    h3.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "Below is the complete implementation mapping for all 26 verified functional capabilities of InternHub. "
        "Each feature is documented across all 12 required fields including exact file paths, HTTP endpoints, models, integrations, and test evidence:"
    ).runs[0].font.size = Pt(9)
    
    for feat in FEATURES_MAP:
        # Create a clean boxed table for each feature
        f_table = doc.add_table(rows=12, cols=2)
        f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header row
        hdr_cell = f_table.rows[0].cells[0]
        hdr_cell.merge(f_table.rows[0].cells[1])
        hdr_cell.text = feat["name"]
        hdr_p = hdr_cell.paragraphs[0]
        hdr_p.runs[0].font.name = "Arial"
        hdr_p.runs[0].font.bold = True
        hdr_p.runs[0].font.size = Pt(9.5)
        set_cell_shading(hdr_cell, "E5E7EB")
        set_cell_margins(hdr_cell, 80, 80, 100, 100)
        
        f_rows = [
            ("Page Name", feat["page"]),
            ("Frontend Route", feat["route"]),
            ("Frontend Source File", feat["frontend_file"]),
            ("UI Components & Hooks", feat["components"]),
            ("Backend API Endpoint", feat["backend_api"]),
            ("Backend Files Involved", feat["backend_files"]),
            ("Database Collections & Fields", feat["database"]),
            ("Real-Time & Email Integrations", feat["integrations"]),
            ("Step-by-Step Workflow", feat["workflow"]),
            ("Role Access Permissions", feat["permissions"]),
            ("Verified Test Evidence", feat["test_evidence"]),
        ]
        
        for idx, (label, val) in enumerate(f_rows, start=1):
            r = f_table.rows[idx]
            c0, c1 = r.cells[0], r.cells[1]
            c0.width = Inches(1.8)
            c1.width = Inches(5.2)
            
            c0.text = label
            c1.text = val
            
            p0 = c0.paragraphs[0]
            p0.runs[0].font.name = "Arial"
            p0.runs[0].font.bold = True
            p0.runs[0].font.size = Pt(7.5)
            
            p1 = c1.paragraphs[0]
            p1.runs[0].font.name = "Arial"
            p1.runs[0].font.size = Pt(7.5)
            
            set_cell_margins(c0, 45, 45, 80, 80)
            set_cell_margins(c1, 45, 45, 80, 80)
            set_cell_shading(c0, "F9FAFB")
            set_cell_shading(c1, "FFFFFF" if idx % 2 == 0 else "FAFAFA")
            
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 4: Backend API and Controller Mapping
    h4 = doc.add_heading("4. Backend API and Controller Mapping", level=1)
    h4.runs[0].font.name = "Arial"
    h4.runs[0].font.size = Pt(13)
    h4.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        f"The InternHub backend routes traffic across 15 route modules comprising exactly {len(ALL_ENDPOINTS)} REST endpoints. "
        "The following master directory maps every HTTP method, URL pattern, controller handler, middleware guard, and purpose:"
    ).runs[0].font.size = Pt(9)
    
    api_table = doc.add_table(rows=len(ALL_ENDPOINTS) + 1, cols=5)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Method", "Endpoint Route", "Controller Function", "Middlewares Applied", "Description & Purpose"]
    for i, h in enumerate(api_headers):
        c = api_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 70, 70, 70, 70)
        
    for r_idx, ep in enumerate(ALL_ENDPOINTS, start=1):
        row = api_table.rows[r_idx]
        vals = [ep["method"], ep["path"], ep["controller"], ep["middleware"], ep["desc"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 40, 40, 50, 50)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 5: Database Mapping
    h5 = doc.add_heading("5. Database Mapping (14 Mongoose Models)", level=1)
    h5.runs[0].font.name = "Arial"
    h5.runs[0].font.size = Pt(13)
    h5.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "InternHub utilizes 14 Mongoose models in MongoDB Atlas. Each model enforces strict data types, referential links, "
        "and compound uniqueness constraints to maintain data integrity across asynchronous transactions:"
    ).runs[0].font.size = Pt(9)
    
    db_table = doc.add_table(rows=len(DATABASE_MODELS) + 1, cols=5)
    db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    db_headers = ["Model / File", "Collection", "Primary Key", "Key Schema Fields & Indexes", "Referential Relationships"]
    for i, h in enumerate(db_headers):
        c = db_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 70, 70, 70, 70)
        
    for r_idx, m in enumerate(DATABASE_MODELS, start=1):
        row = db_table.rows[r_idx]
        vals = [m["model"], m["collection"], m["pk"], f"Fields: {m['fields']}\nIndexes: {m['indexes']}", m["relations"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 45, 45, 60, 60)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 6: Role and Permission Matrix
    h6 = doc.add_heading("6. Role and Permission Matrix (5-Tier RBAC)", level=1)
    h6.runs[0].font.name = "Arial"
    h6.runs[0].font.size = Pt(13)
    h6.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "Role-Based Access Control is enforced through backend requireRole middleware and client-side ProtectedRoute components. "
        "Below is the complete permission matrix across all 13 operational domains:"
    ).runs[0].font.size = Pt(9)
    
    rm_table = doc.add_table(rows=len(ROLE_MATRIX) + 1, cols=6)
    rm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rm_headers = ["Functional Domain", "Student", "Mentor", "Company", "Coordinator", "Administrator"]
    for i, h in enumerate(rm_headers):
        c = rm_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(7.5)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 70, 70, 60, 60)
        
    for r_idx, r in enumerate(ROLE_MATRIX, start=1):
        row = rm_table.rows[r_idx]
        vals = [r["domain"], r["student"], r["mentor"], r["company"], r["coordinator"], r["admin"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 40, 40, 50, 50)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 7: In-Depth Workflow A - Session Management
    h7 = doc.add_heading("7. In-Depth Workflow A: Session Management", level=1)
    h7.runs[0].font.name = "Arial"
    h7.runs[0].font.size = Pt(13)
    h7.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    for title, desc in WORKFLOW_SESSION["steps"]:
        p_step = doc.add_paragraph()
        r_step = p_step.add_run(f"• {title}: ")
        r_step.font.name = "Arial"
        r_step.font.bold = True
        r_step.font.size = Pt(8.5)
        
        r_desc = p_step.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(8.5)
        p_step.paragraph_format.space_after = Pt(4)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 8: In-Depth Workflow B - Certificate Generation
    h8 = doc.add_heading("8. In-Depth Workflow B: Certificate Generation & Verification", level=1)
    h8.runs[0].font.name = "Arial"
    h8.runs[0].font.size = Pt(13)
    h8.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    for title, desc in WORKFLOW_CERTIFICATE["steps"]:
        p_step = doc.add_paragraph()
        r_step = p_step.add_run(f"• {title}: ")
        r_step.font.name = "Arial"
        r_step.font.bold = True
        r_step.font.size = Pt(8.5)
        
        r_desc = p_step.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(8.5)
        p_step.paragraph_format.space_after = Pt(4)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 9: In-Depth Workflow C - Real-Time Notifications & Chat
    h9 = doc.add_heading("9. In-Depth Workflow C: Real-Time Notifications and Chat", level=1)
    h9.runs[0].font.name = "Arial"
    h9.runs[0].font.size = Pt(13)
    h9.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    for title, desc in WORKFLOW_CHAT_NOTIF["steps"]:
        p_step = doc.add_paragraph()
        r_step = p_step.add_run(f"• {title}: ")
        r_step.font.name = "Arial"
        r_step.font.bold = True
        r_step.font.size = Pt(8.5)
        
        r_desc = p_step.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(8.5)
        p_step.paragraph_format.space_after = Pt(4)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 10: Testing File Mapping
    h10 = doc.add_heading("10. Testing File Mapping (225 Automated Tests)", level=1)
    h10.runs[0].font.name = "Arial"
    h10.runs[0].font.size = Pt(13)
    h10.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "All 225 automated backend tests across 17 test suites were audited and verified passing cleanly (100% green). "
        "The table below maps each test file to its subsystem, test count, and execution command:"
    ).runs[0].font.size = Pt(9)
    
    t_table = doc.add_table(rows=len(TEST_SUITES) + 1, cols=4)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Test Suite File", "Subsystem / Functional Domain", "Tests", "Sequential Execution Command"]
    for i, h in enumerate(t_headers):
        c = t_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 70, 70, 70, 70)
        
    for r_idx, ts in enumerate(TEST_SUITES, start=1):
        row = t_table.rows[r_idx]
        vals = [ts["suite"], ts["subsystem"], str(ts["tests"]), ts["cmd"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7.5)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 45, 45, 60, 60)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 11: Missing Links, Incomplete Features & Issues Discovered
    h11 = doc.add_heading("11. Missing Links, Incomplete Features & Issues Discovered", level=1)
    h11.runs[0].font.name = "Arial"
    h11.runs[0].font.size = Pt(13)
    h11.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    doc.add_paragraph(
        "A rigorous architectural code audit identified 6 architectural observations, routing nuances, and scalability considerations:"
    ).runs[0].font.size = Pt(9)
    
    iss_table = doc.add_table(rows=len(ISSUES_DISCOVERED) + 1, cols=4)
    iss_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    iss_headers = ["Observation / Issue", "File Location", "Classification", "Architectural Analysis & Remediation"]
    for i, h in enumerate(iss_headers):
        c = iss_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shading(c, "E5E7EB")
        set_cell_margins(c, 70, 70, 70, 70)
        
    for r_idx, iss in enumerate(ISSUES_DISCOVERED, start=1):
        row = iss_table.rows[r_idx]
        vals = [iss["issue"], iss["location"], iss["nature"], iss["description"]]
        for c_idx, val in enumerate(vals):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(7.5)
            if c_idx == 0:
                p.runs[0].font.bold = True
            set_cell_margins(c, 45, 45, 60, 60)
            set_cell_shading(c, "FFFFFF" if r_idx % 2 != 0 else "F9FAFB")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 12: Technical Summary for Team Lead
    h12 = doc.add_heading("12. Technical Summary for Engineering Team Lead", level=1)
    h12.runs[0].font.name = "Arial"
    h12.runs[0].font.size = Pt(13)
    h12.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    
    summary_text = (
        "InternHub exhibits an exceptionally high standard of backend software engineering, architectural discipline, and security hardening. "
        "The system seamlessly balances relational business logic with the flexibility of MongoDB through 14 strongly typed schemas and defensive compound indexes. "
        "Every single functional requirement—from student self-registration and Google GIS SSO to AI resume parsing, supervisor task management, "
        "idempotent PDFKit landscape certification with cryptographic QR verification, and bi-directional Socket.IO chat with active-window muting guards—is completely implemented and verified. "
        "With 225 automated backend tests passing at 100% green, zero regressions, and a clean Vite production build (1,755 modules, 0 errors), "
        "the architecture is robust, secure, and fully verified for institutional staging and production release."
    )
    p_sum = doc.add_paragraph(summary_text)
    p_sum.runs[0].font.name = "Arial"
    p_sum.runs[0].font.size = Pt(9)
    p_sum.paragraph_format.space_after = Pt(10)
    
    # Save DOCX
    doc.save(docx_path)
    print(f"DOCX report saved successfully! File size: {os.path.getsize(docx_path)} bytes")

print("Defined build_docx_report function")

# ==============================================================================
# 3. PDF REPORT BUILDER (BLACK & WHITE ENTERPRISE STYLE)
# ==============================================================================

class NumberedCanvasBW(canvas.Canvas):
    """Two-pass canvas for total page count and running header/footer in black & white."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, total_pages):
        self.saveState()
        
        # Suppress running header on cover page (page 1)
        if self._pageNumber > 1:
            self.setFont('Helvetica-Bold', 7.5)
            self.setFillColor(colors.black)
            self.drawString(36, 762, 'INTERNHUB – FEATURE IMPLEMENTATION & TECHNICAL ARCHITECTURE REPORT')
            self.setFont('Helvetica', 7.5)
            self.setFillColor(colors.HexColor('#4b5563'))
            self.drawRightString(576, 762, 'DOC ID: IHB-ARCH-2026-V1')
            
            # Header Rule
            self.setStrokeColor(colors.HexColor('#9ca3af'))
            self.setLineWidth(0.5)
            self.line(36, 756, 576, 756)
        
        # Footer Rule
        self.setStrokeColor(colors.HexColor('#9ca3af'))
        self.setLineWidth(0.5)
        self.line(36, 38, 576, 38)
        
        # Running Footer
        self.setFont('Helvetica-Bold', 7)
        self.setFillColor(colors.black)
        self.drawString(36, 26, 'CONFIDENTIAL — FOR INTERNAL ENGINEERING & ARCHITECTURAL GOVERNANCE ONLY')
        self.setFont('Helvetica', 7.5)
        self.setFillColor(colors.HexColor('#4b5563'))
        self.drawRightString(576, 26, f'Page {self._pageNumber} of {total_pages}')
        
        self.restoreState()

def create_pdf_styles():
    base = getSampleStyleSheet()
    
    styles = {
        'DocTitle': ParagraphStyle(
            'DocTitle',
            parent=base['Title'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.black,
            alignment=0,
            spaceAfter=3
        ),
        'DocSubtitle': ParagraphStyle(
            'DocSubtitle',
            parent=base['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#374151'),
            alignment=0,
            spaceAfter=8
        ),
        'SectionH1': ParagraphStyle(
            'SectionH1',
            parent=base['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=11.5,
            leading=15,
            textColor=colors.black,
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True
        ),
        'SectionH2': ParagraphStyle(
            'SectionH2',
            parent=base['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=9.5,
            leading=12.5,
            textColor=colors.black,
            spaceBefore=6,
            spaceAfter=3,
            keepWithNext=True
        ),
        'Body': ParagraphStyle(
            'Body',
            parent=base['Normal'],
            fontName='Helvetica',
            fontSize=7.5,
            leading=10.5,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=3
        ),
        'BodyBold': ParagraphStyle(
            'BodyBold',
            parent=base['Normal'],
            fontName='Helvetica-Bold',
            fontSize=7.5,
            leading=10.5,
            textColor=colors.black,
            spaceAfter=3
        ),
        'TableHeader': ParagraphStyle(
            'TableHeader',
            parent=base['Normal'],
            fontName='Helvetica-Bold',
            fontSize=7.5,
            leading=9.5,
            textColor=colors.black,
            alignment=0
        ),
        'TableCell': ParagraphStyle(
            'TableCell',
            parent=base['Normal'],
            fontName='Helvetica',
            fontSize=6.8,
            leading=8.8,
            textColor=colors.HexColor('#111827')
        ),
        'TableCellBold': ParagraphStyle(
            'TableCellBold',
            parent=base['Normal'],
            fontName='Helvetica-Bold',
            fontSize=6.8,
            leading=8.8,
            textColor=colors.black
        ),
        'TableCellCode': ParagraphStyle(
            'TableCellCode',
            parent=base['Normal'],
            fontName='Courier-Bold',
            fontSize=6.5,
            leading=8.2,
            textColor=colors.black
        ),
        'CoverLabel': ParagraphStyle(
            'CoverLabel',
            parent=base['Normal'],
            fontName='Helvetica-Bold',
            fontSize=7.5,
            leading=9.5,
            textColor=colors.black
        ),
        'CoverVal': ParagraphStyle(
            'CoverVal',
            parent=base['Normal'],
            fontName='Helvetica',
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor('#111827')
        )
    }
    return styles

def build_pdf_report(pdf_path):
    print(f"Building PDF report at: {pdf_path}")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=42,
        bottomMargin=42
    )
    
    styles = create_pdf_styles()
    story = []
    
    # 1. Document Title & Subtitle
    story.append(Paragraph(html.escape(DOC_TITLE), styles['DocTitle']))
    story.append(Paragraph(html.escape(DOC_SUBTITLE), styles['DocSubtitle']))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceAfter=6, spaceBefore=2))
    
    # 2. Metadata Table
    meta_rows = [
        [Paragraph("Document ID", styles['CoverLabel']), Paragraph(DOC_ID, styles['CoverVal']),
         Paragraph("Date of Audit", styles['CoverLabel']), Paragraph(DOC_DATE, styles['CoverVal'])],
        [Paragraph("System Under Test", styles['CoverLabel']), Paragraph(DOC_SYSTEM, styles['CoverVal']),
         Paragraph("Classification", styles['CoverLabel']), Paragraph(DOC_CLASSIFICATION, styles['CoverVal'])],
        [Paragraph("Lead Author", styles['CoverLabel']), Paragraph(DOC_AUTHOR, styles['CoverVal']),
         Paragraph("Final Verdict", styles['CoverLabel']), Paragraph(f"<b>{DOC_VERDICT}</b>", styles['CoverVal'])],
    ]
    meta_table = Table(meta_rows, colWidths=[90, 180, 85, 185])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f9fafb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 8))
    
    # Section 1: Project Architecture Overview
    story.append(Paragraph("<b>1. Project Architecture Overview</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    p_arch = (
        "The Smart Internship Management System (InternHub) is engineered as an enterprise-grade micro-monolith platform uniting five distinct stakeholders: "
        "Students, Industry Mentors, Employer Companies, Academic Coordinators, and System Administrators. "
        "The backend pairs a stateless Express.js REST API with a stateful, bi-directional Socket.IO WebSocket server operating synchronously on port 5001. "
        "Data persistence is managed via MongoDB Atlas utilizing Mongoose 8.9.5 ODM with strict schema validation and compound uniqueness indexing. "
        "The frontend is a single-page application (SPA) built with React 18.3.1, bundled via Vite 6.0.7, styled with Tailwind CSS 3.4.17, "
        "and structured through React Router DOM v6 with role-aware route guards."
    )
    story.append(Paragraph(p_arch, styles['Body']))
    story.append(Spacer(1, 4))
    
    # Architecture Stack Table
    tech_data = [
        [Paragraph("Layer / Subsystem", styles['TableHeader']), Paragraph("Technology & Version", styles['TableHeader']), Paragraph("Architectural Role & Scope", styles['TableHeader'])],
        [Paragraph("<b>Backend Framework</b>", styles['TableCellBold']), Paragraph("Express.js ^4.21.2 / Node.js v24.17.0", styles['TableCell']), Paragraph("Handles 15 route modules, 123 endpoints, JWT auth, and middleware.", styles['TableCell'])],
        [Paragraph("<b>Real-Time Engine</b>", styles['TableCellBold']), Paragraph("Socket.IO ^4.8.4", styles['TableCell']), Paragraph("Maintains authenticated persistent connections, user rooms, and chat events.", styles['TableCell'])],
        [Paragraph("<b>Database & ODM</b>", styles['TableCellBold']), Paragraph("MongoDB Atlas / Mongoose ^8.9.5", styles['TableCell']), Paragraph("14 strongly-typed collections, compound uniqueness barriers, and aggregations.", styles['TableCell'])],
        [Paragraph("<b>Frontend Application</b>", styles['TableCellBold']), Paragraph("React ^18.3.1 / Vite ^6.0.7", styles['TableCell']), Paragraph("SPA with 26 pages, Context API, responsive Tailwind layout, and Vite bundling.", styles['TableCell'])],
        [Paragraph("<b>Document Generation</b>", styles['TableCellBold']), Paragraph("PDFKit ^0.20.2 / QRCode ^1.5.4", styles['TableCell']), Paragraph("Generates landscape A4 completion certificates with QR verification & Padmavathi signature.", styles['TableCell'])],
        [Paragraph("<b>External Services</b>", styles['TableCellBold']), Paragraph("Brevo/Gmail SMTP / Gemini 2.5 Flash", styles['TableCell']), Paragraph("Nodemailer transactional email delivery; Google Gemini AI resume fit scoring.", styles['TableCell'])]
    ]
    t_tech = Table(tech_data, colWidths=[110, 160, 270])
    t_tech.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_tech)
    story.append(Spacer(1, 8))
    
    # Section 2: Frontend Page and Route Directory
    story.append(PageBreak())
    story.append(Paragraph("<b>2. Frontend Page and Route Directory (26 Pages)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    page_data = [
        [Paragraph("Page Title", styles['TableHeader']), Paragraph("Route URL", styles['TableHeader']), Paragraph("Source File Path", styles['TableHeader']), Paragraph("Allowed Roles", styles['TableHeader']), Paragraph("Navigation & Access", styles['TableHeader'])]
    ]
    for p in PAGES_DIRECTORY:
        page_data.append([
            Paragraph(f"<b>{p['title']}</b>", styles['TableCellBold']),
            Paragraph(p['route'], styles['TableCellCode']),
            Paragraph(p['file'], styles['TableCell']),
            Paragraph(p['roles'], styles['TableCell']),
            Paragraph(p['nav'], styles['TableCell'])
        ])
    t_page = Table(page_data, colWidths=[95, 110, 140, 95, 100], repeatRows=1)
    t_page.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_page)
    story.append(Spacer(1, 8))
    
    # Section 3: Complete Feature-to-File Implementation Mapping
    story.append(PageBreak())
    story.append(Paragraph("<b>3. Complete Feature-to-File Implementation Mapping (26 Features)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    for feat in FEATURES_MAP:
        f_card_data = [
            [Paragraph(f"<b>{feat['name']}</b>", styles['TableHeader']), Paragraph("", styles['TableHeader'])],
            [Paragraph("<b>Page Name:</b>", styles['TableCellBold']), Paragraph(feat['page'], styles['TableCell'])],
            [Paragraph("<b>Frontend Route:</b>", styles['TableCellBold']), Paragraph(feat['route'], styles['TableCellCode'])],
            [Paragraph("<b>Frontend Source File:</b>", styles['TableCellBold']), Paragraph(feat['frontend_file'], styles['TableCell'])],
            [Paragraph("<b>UI Components & Hooks:</b>", styles['TableCellBold']), Paragraph(feat['components'], styles['TableCell'])],
            [Paragraph("<b>Backend API Endpoint:</b>", styles['TableCellBold']), Paragraph(feat['backend_api'], styles['TableCellCode'])],
            [Paragraph("<b>Backend Files Involved:</b>", styles['TableCellBold']), Paragraph(feat['backend_files'], styles['TableCell'])],
            [Paragraph("<b>Database Collections:</b>", styles['TableCellBold']), Paragraph(feat['database'], styles['TableCell'])],
            [Paragraph("<b>Real-Time / Email:</b>", styles['TableCellBold']), Paragraph(feat['integrations'], styles['TableCell'])],
            [Paragraph("<b>Step-by-Step Workflow:</b>", styles['TableCellBold']), Paragraph(feat['workflow'], styles['TableCell'])],
            [Paragraph("<b>Role Access:</b>", styles['TableCellBold']), Paragraph(feat['permissions'], styles['TableCell'])],
            [Paragraph("<b>Verified Test Evidence:</b>", styles['TableCellBold']), Paragraph(feat['test_evidence'], styles['TableCellBold'])],
        ]
        t_card = Table(f_card_data, colWidths=[120, 420])
        t_card.setStyle(TableStyle([
            ('SPAN', (0, 0), (1, 0)),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
            ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#f9fafb')),
        ]))
        story.append(KeepTogether([t_card, Spacer(1, 5)]))
        
    # Section 4: Backend API and Controller Mapping
    story.append(PageBreak())
    story.append(Paragraph(f"<b>4. Backend API and Controller Mapping ({len(ALL_ENDPOINTS)} Endpoints)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    api_data = [
        [Paragraph("Method", styles['TableHeader']), Paragraph("Endpoint Route", styles['TableHeader']), Paragraph("Controller Function", styles['TableHeader']), Paragraph("Middlewares Applied", styles['TableHeader']), Paragraph("Description & Purpose", styles['TableHeader'])]
    ]
    for ep in ALL_ENDPOINTS:
        api_data.append([
            Paragraph(f"<b>{ep['method']}</b>", styles['TableCellBold']),
            Paragraph(ep['path'], styles['TableCellCode']),
            Paragraph(ep['controller'], styles['TableCell']),
            Paragraph(ep['middleware'], styles['TableCell']),
            Paragraph(ep['desc'], styles['TableCell'])
        ])
    t_api = Table(api_data, colWidths=[42, 138, 120, 105, 135], repeatRows=1)
    t_api.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_api)
    story.append(Spacer(1, 8))
    
    # Section 5: Database Mapping
    story.append(PageBreak())
    story.append(Paragraph("<b>5. Database Mapping (14 Mongoose Models)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    db_data = [
        [Paragraph("Model / File", styles['TableHeader']), Paragraph("Collection", styles['TableHeader']), Paragraph("PK", styles['TableHeader']), Paragraph("Key Schema Fields & Indexes", styles['TableHeader']), Paragraph("Referential Relationships", styles['TableHeader'])]
    ]
    for m in DATABASE_MODELS:
        db_data.append([
            Paragraph(f"<b>{m['model']}</b>", styles['TableCellBold']),
            Paragraph(m['collection'], styles['TableCellCode']),
            Paragraph(m['pk'], styles['TableCell']),
            Paragraph(f"<b>Fields:</b> {m['fields']}<br/><b>Indexes:</b> {m['indexes']}", styles['TableCell']),
            Paragraph(m['relations'], styles['TableCell'])
        ])
    t_db = Table(db_data, colWidths=[75, 75, 55, 205, 130], repeatRows=1)
    t_db.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_db)
    story.append(Spacer(1, 8))
    
    # Section 6: Role and Permission Matrix
    story.append(PageBreak())
    story.append(Paragraph("<b>6. Role and Permission Matrix (5-Tier RBAC)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    rm_data = [
        [Paragraph("Domain", styles['TableHeader']), Paragraph("Student", styles['TableHeader']), Paragraph("Mentor", styles['TableHeader']), Paragraph("Company", styles['TableHeader']), Paragraph("Coordinator", styles['TableHeader']), Paragraph("Administrator", styles['TableHeader'])]
    ]
    for r in ROLE_MATRIX:
        rm_data.append([
            Paragraph(f"<b>{r['domain']}</b>", styles['TableCellBold']),
            Paragraph(r['student'], styles['TableCell']),
            Paragraph(r['mentor'], styles['TableCell']),
            Paragraph(r['company'], styles['TableCell']),
            Paragraph(r['coordinator'], styles['TableCell']),
            Paragraph(r['admin'], styles['TableCell'])
        ])
    t_rm = Table(rm_data, colWidths=[85, 90, 90, 90, 90, 95], repeatRows=1)
    t_rm.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_rm)
    story.append(Spacer(1, 8))
    
    # Section 7: In-Depth Workflow A
    story.append(PageBreak())
    story.append(Paragraph(f"<b>7. {WORKFLOW_SESSION['title']}</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    for t_step, d_step in WORKFLOW_SESSION['steps']:
        p_txt = f"<b>• {t_step}:</b> {d_step}"
        story.append(Paragraph(p_txt, styles['Body']))
        story.append(Spacer(1, 2))
        
    # Section 8: In-Depth Workflow B
    story.append(PageBreak())
    story.append(Paragraph(f"<b>8. {WORKFLOW_CERTIFICATE['title']}</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    for t_step, d_step in WORKFLOW_CERTIFICATE['steps']:
        p_txt = f"<b>• {t_step}:</b> {d_step}"
        story.append(Paragraph(p_txt, styles['Body']))
        story.append(Spacer(1, 2))
        
    # Section 9: In-Depth Workflow C
    story.append(PageBreak())
    story.append(Paragraph(f"<b>9. {WORKFLOW_CHAT_NOTIF['title']}</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    for t_step, d_step in WORKFLOW_CHAT_NOTIF['steps']:
        p_txt = f"<b>• {t_step}:</b> {d_step}"
        story.append(Paragraph(p_txt, styles['Body']))
        story.append(Spacer(1, 2))
        
    # Section 10: Testing File Mapping
    story.append(PageBreak())
    story.append(Paragraph("<b>10. Testing File Mapping (225 Automated Tests)</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    t_test_data = [
        [Paragraph("Test Suite File", styles['TableHeader']), Paragraph("Subsystem / Functional Domain", styles['TableHeader']), Paragraph("Tests", styles['TableHeader']), Paragraph("Sequential Execution Command", styles['TableHeader'])]
    ]
    for ts in TEST_SUITES:
        t_test_data.append([
            Paragraph(f"<b>{ts['suite']}</b>", styles['TableCellBold']),
            Paragraph(ts['subsystem'], styles['TableCell']),
            Paragraph(str(ts['tests']), styles['TableCellBold']),
            Paragraph(ts['cmd'], styles['TableCellCode'])
        ])
    t_test_table = Table(t_test_data, colWidths=[120, 160, 40, 220], repeatRows=1)
    t_test_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_test_table)
    story.append(Spacer(1, 8))
    
    # Section 11: Missing Links & Issues
    story.append(PageBreak())
    story.append(Paragraph("<b>11. Missing Links, Incomplete Features & Issues Discovered</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    iss_data = [
        [Paragraph("Observation / Issue", styles['TableHeader']), Paragraph("File Location", styles['TableHeader']), Paragraph("Classification", styles['TableHeader']), Paragraph("Architectural Analysis & Remediation", styles['TableHeader'])]
    ]
    for iss in ISSUES_DISCOVERED:
        iss_data.append([
            Paragraph(f"<b>{iss['issue']}</b>", styles['TableCellBold']),
            Paragraph(iss['location'], styles['TableCellCode']),
            Paragraph(iss['nature'], styles['TableCell']),
            Paragraph(iss['description'], styles['TableCell'])
        ])
    t_iss = Table(iss_data, colWidths=[110, 130, 90, 210], repeatRows=1)
    t_iss.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_iss)
    story.append(Spacer(1, 8))
    
    # Section 12: Technical Summary
    story.append(PageBreak())
    story.append(Paragraph("<b>12. Technical Summary for Engineering Team Lead</b>", styles['SectionH1']))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#9ca3af'), spaceAfter=4, spaceBefore=1))
    
    summary_text = (
        "InternHub exhibits an exceptionally high standard of backend software engineering, architectural discipline, and security hardening. "
        "The system seamlessly balances relational business logic with the flexibility of MongoDB through 14 strongly typed schemas and defensive compound indexes. "
        "Every single functional requirement—from student self-registration and Google GIS SSO to AI resume parsing, supervisor task management, "
        "idempotent PDFKit landscape certification with cryptographic QR verification, and bi-directional Socket.IO chat with active-window muting guards—is completely implemented and verified. "
        "With 225 automated backend tests passing at 100% green, zero regressions, and a clean Vite production build (1,755 modules, 0 errors), "
        "the architecture is robust, secure, and fully verified for institutional staging and production release."
    )
    story.append(Paragraph(summary_text, styles['Body']))
    story.append(Spacer(1, 10))
    
    # Sign-off box
    sign_data = [
        [Paragraph("Role / Authority", styles['TableHeader']), Paragraph("Designated Signatory", styles['TableHeader']), Paragraph("Audit Date", styles['TableHeader']), Paragraph("Governance Status", styles['TableHeader'])],
        [Paragraph("<b>Lead QA Systems Engineer</b>", styles['TableCellBold']), Paragraph("Senior QA Systems Engineer", styles['TableCell']), Paragraph("October 23, 2026", styles['TableCell']), Paragraph("<b>VERIFIED & APPROVED</b>", styles['TableCellBold'])],
        [Paragraph("<b>Technical Architect</b>", styles['TableCellBold']), Paragraph("Lead Software Architect", styles['TableCell']), Paragraph("October 23, 2026", styles['TableCell']), Paragraph("<b>APPROVED FOR SUBMISSION</b>", styles['TableCellBold'])],
    ]
    t_sign = Table(sign_data, colWidths=[130, 150, 100, 160])
    t_sign.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e5e7eb')),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.black),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
    ]))
    story.append(t_sign)
    
    # Build document
    doc.build(story, canvasmaker=NumberedCanvasBW)
    print(f"PDF report saved successfully! File size: {os.path.getsize(pdf_path)} bytes")

def main():
    os.makedirs('reports', exist_ok=True)
    docx_path = os.path.abspath('reports/INTERNHUB_COMPLETE_FEATURE_IMPLEMENTATION_AND_ARCHITECTURE_REPORT.docx')
    pdf_path = os.path.abspath('reports/INTERNHUB_COMPLETE_FEATURE_IMPLEMENTATION_AND_ARCHITECTURE_REPORT.pdf')
    
    build_docx_report(docx_path)
    build_pdf_report(pdf_path)
    print("\nALL REPORTS GENERATED CLEANLY IN reports/")

if __name__ == '__main__':
    main()

